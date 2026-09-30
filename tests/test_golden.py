"""Regression ("golden") tests: a refactor must not move a valuation.

The whole pipeline runs on the synthetic snapshot from conftest, so this needs no network and no
Yahoo data. It pins the numbers an analyst would quote from the deck, checks the arithmetic that
ties them together, and asserts that all four deliverables are produced.

If a change moves a number on purpose, update PINNED here and say so in the commit message. The
Excel-reproduces-Python check needs Excel itself and is skipped where it is not installed.
"""
from __future__ import annotations

import json
import sys

import pytest
from openpyxl import load_workbook

import eqr.pipeline as pl
from eqr.config import CompanyConfig, PeerConfig
from eqr.retrieve import DiskCache, Extras
from tests.conftest import make_snapshot

# What the synthetic case is worth. Every figure traces to a statement line in conftest.make_snapshot.
PINNED = {
    "value_per_share": 25.385103, "enterprise_value": 2738.510288,
    "sum_pv_fcf": 830.130589, "pv_terminal": 1908.379699,
    "wacc": 0.091198, "cost_of_equity": 0.095500, "beta_used": 1.0,
    "target_price": 27.500000, "fair_value": 25.385103, "multiples_value": 43.969453,
    "upside": -0.450000,
}
PINNED_RATING = "SELL"
PINNED_SCENARIOS = {"Bear": 18.528996, "Base": 25.385103}


@pytest.fixture
def case(monkeypatch):
    """The synthetic company plus two peers, wired so no call leaves the machine."""
    company = make_snapshot()
    snapshots = {
        "TEST.OL": company,
        "P1.OL": make_snapshot("P1.OL", price=40.0, trailing_pe=15.0, forward_pe=12.0),
        "P2.OL": make_snapshot("P2.OL", price=60.0, ebitda_margin=0.35, trailing_pe=21.0, forward_pe=16.0),
    }
    monkeypatch.setattr(pl, "fetch_snapshot", lambda symbol, cache, **kw: snapshots[symbol])
    monkeypatch.setattr(pl, "fetch_prices", lambda symbol, cache, **kw: company.prices)
    monkeypatch.setattr(pl, "fetch_fx", lambda a, b, cache=None: 1.0)
    monkeypatch.setattr(pl, "fetch_extras", lambda symbol, cache, **kw: Extras(ticker=symbol))
    cfg = CompanyConfig(ticker="TEST.OL", name="Test ASA", index="^TEST", units_label="NOKm")
    cfg.peers = [PeerConfig("P1.OL", "Peer One", "Nordic"), PeerConfig("P2.OL", "Peer Two", "Nordic")]
    # Valued at the fiscal year-end: a date-free, textbook DCF, so the pinned numbers never drift with the calendar.
    # The stub period (valuation date after the year-end) is pinned separately in test_valuation_date.py.
    cfg.assumptions.valuation_date = "2025-12-31"
    return pl.run_analysis(cfg, DiskCache(enabled=False), progress=lambda m: None)


def headline(result) -> dict:
    d = json.loads(result.to_json())
    dcf, wacc, rec = d["dcf"], d["wacc"], d["recommendation"]
    return {
        "value_per_share": dcf["value_per_share"], "enterprise_value": dcf["enterprise_value"],
        "sum_pv_fcf": dcf["sum_pv_fcf"], "pv_terminal": dcf["pv_terminal"],
        "wacc": wacc["wacc"], "cost_of_equity": wacc["cost_of_equity"], "beta_used": wacc["beta_used"],
        "target_price": rec["target_price"], "fair_value": rec["fair_value"],
        "multiples_value": rec["multiples_value"], "upside": rec["upside"],
    }


def test_the_synthetic_case_is_valued_exactly_as_before(case):
    assert headline(case) == pytest.approx(PINNED, abs=2e-6)
    d = json.loads(case.to_json())
    assert d["recommendation"]["rating"] == PINNED_RATING
    scenarios = {s["name"]: s["value_per_share"] for s in d["scenarios"]}
    assert {k: scenarios[k] for k in PINNED_SCENARIOS} == pytest.approx(PINNED_SCENARIOS, abs=2e-6)


def test_the_valuation_adds_up(case):
    """The pinned numbers are not just stable, they are consistent: this is the arithmetic a reader checks."""
    d = json.loads(case.to_json())
    dcf, rec = d["dcf"], d["recommendation"]
    assert dcf["enterprise_value"] == pytest.approx(dcf["sum_pv_fcf"] + dcf["pv_terminal"])
    assert dcf["equity_value"] == pytest.approx(dcf["enterprise_value"] - dcf["net_debt"] - dcf["minorities"])
    assert dcf["value_per_share"] == pytest.approx(dcf["equity_value"] / dcf["shares"])
    assert rec["upside"] == pytest.approx(rec["target_price"] / rec["price"] - 1)
    # the target is the fair value rolled forward one year at the cost of equity, less the dividend
    horizon = rec["horizon_months"] / 12
    rolled = rec["fair_value"] * (1 + rec["cost_of_equity"]) ** horizon - rec.get("dps", 0.0) * horizon
    assert rec["target_price"] == pytest.approx(rolled, rel=0.02)
    weighted = sum(s["probability"] * s["value_per_share"] for s in d["scenarios"])
    assert d["scenario_weighted_value"] == pytest.approx(weighted, rel=1e-9)
    assert sum(s["probability"] for s in d["scenarios"]) == pytest.approx(1.0)


def test_every_deliverable_is_produced_and_carries_the_numbers(case, tmp_path):
    from eqr.create.dashboard import build_dashboard
    from eqr.create.deck import build_deck
    from eqr.create.excel import build_excel
    from eqr.narrative import generate_narrative

    case.narrative, case.narrative_mode = generate_narrative(case, mode="rules")
    deck = build_deck(case, tmp_path / "deck.pptx", charts_dir=tmp_path / "charts")
    excel = build_excel(case, tmp_path / "model.xlsx")
    dashboard = build_dashboard(case, tmp_path / "dashboard.html")
    payload = tmp_path / "analysis.json"
    payload.write_text(case.to_json(), encoding="utf-8")

    for path in (deck, excel, dashboard, payload):
        assert path.exists() and path.stat().st_size > 5_000, path.name

    from pptx import Presentation

    slides = Presentation(deck).slides
    assert len(slides) >= 6                                    # cover, team, the four fixed headings
    headings = " ".join(shape.text_frame.text for slide in slides for shape in slide.shapes
                        if shape.has_text_frame).lower()
    for required in ("company overview", "market overview", "financials", "valuation"):
        assert required in headings, required

    book = load_workbook(excel)                                # formulas, not values: openpyxl does not calculate
    assert {"Inputs", "DCF", "Checks"} <= set(book.sheetnames)
    assert str(book["DCF"]["B21"].value).startswith("=")       # fair value is a live formula, not a pasted number

    html = dashboard.read_text(encoding="utf-8")
    assert f"{case.recommendation.target_price:.2f}" in html or f"{case.recommendation.target_price:.1f}" in html


@pytest.mark.skipif(not sys.platform.startswith("win"), reason="the workbook is recalculated through Excel COM")
def test_excel_reproduces_python(case, tmp_path):
    """The model's own formulas must land on the same fair value, and every integrity check must read TRUE.

    openpyxl writes formulas but cannot evaluate them, so this asks Excel to recalculate and save, then
    reads the cached values back. Skipped where Excel is not available (CI runs on Linux).
    """
    from eqr.create.excel import build_excel, recalc_with_excel

    case.narrative, case.narrative_mode = "", "rules"
    path = build_excel(case, tmp_path / "model.xlsx")
    try:
        report = recalc_with_excel(path)
    except Exception as exc:                                   # Excel present but not usable (no licence, locked)
        pytest.skip(f"Excel could not recalculate the workbook: {exc}")
    if report.get("status") == "skipped":
        pytest.skip("Excel COM unavailable")
    assert report["status"] == "success", report["errors"]     # no #REF!, #DIV/0! or #VALUE! anywhere

    book = load_workbook(path, data_only=True)                 # now the cached values exist
    assert book["DCF"]["B21"].value == pytest.approx(case.dcf.value_per_share, rel=2e-4)
    assert book["DCF"]["B30"].value == pytest.approx(case.recommendation.target_price, rel=2e-4)
    checks = book["Checks"]
    failed = [checks.cell(row=r, column=1).value for r in range(5, checks.max_row + 1)
              if checks.cell(row=r, column=2).value is False]
    assert not failed, f"integrity checks that came back FALSE: {failed}"
