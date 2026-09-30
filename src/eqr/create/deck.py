"""PowerPoint deck: cover, team, and the four case slides.

Layout mirrors a Nordic sell-side case template (16:9, 13.33 x 7.5 in): section
tag, serif title, one-line subtitle, numbered section boxes with grey panels,
sources footer and page number. Charts are rendered with matplotlib; tables are
native (editable) PowerPoint tables.
"""
from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..analyze import MULTIPLE_LABELS
from . import charts as C
from .style import (
    ACCENT_BLUE,
    BLACK,
    DARK_GREY,
    FONT_BODY,
    FONT_TITLE,
    GREEN,
    GREY,
    LIGHT_BLUE,
    LIGHT_GREY,
    MID_GREY,
    NAVY,
    NAVY_TEXT,
    PALE_BLUE,
    WHITE,
    fmt_leverage,
    fmt_mult,
    fmt_num,
    fmt_pct,
    fmt_price,
)

SW, SH = 13.333, 7.5
LM = 0.47
CW = SW - 2 * LM
FOOTER_Y = 6.93
SECTION_TAGS = ["#1 The equity story", "#2 Market overview and peer landscape", "#3 Historical analysis and forecasts",
                "#4 Valuation and concluding thoughts"]


# =============================================================================
# low-level helpers
# =============================================================================
def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def _runs(text: str) -> list[tuple[str, bool]]:
    """Split '**bold** rest' markup into (text, bold) runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    out = []
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            out.append((p[2:-2], True))
        else:
            out.append((p, False))
    return out


def _style_run(run, size: float, bold: bool, color: str, font: str, italic: bool = False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = _rgb(color)


def add_text(slide, x, y, w, h, text, *, size=10.0, bold=False, italic=False, color=NAVY_TEXT, font=FONT_BODY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.03, line_spacing=None, wrap=True, markup=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        runs = _runs(str(line)) if markup else [(str(line), False)]
        for t, b in runs:
            r = p.add_run()
            r.text = t
            _style_run(r, size, bold or b, color, font, italic)
    return box


def _set_bullet(paragraph, char="▪", indent_in=0.14, color=NAVY):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": color})
    buClr.append(srgb)
    pPr.append(buClr)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    pPr.append(buFont)
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buChar)


def add_bullets(slide, x, y, w, h, items: Sequence[str], *, size=8.5, color=BLACK, space_after=4, margin=0.06,
                anchor=MSO_ANCHOR.TOP, bullet_color=NAVY, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        _set_bullet(p, color=bullet_color)
        for t, b in _runs(str(item)):
            r = p.add_run()
            r.text = t
            _style_run(r, size, b, color, FONT_BODY)
    return box


def add_rect(slide, x, y, w, h, *, fill: str | None = LIGHT_GREY, line: str | None = None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


def add_line(slide, x1, y1, x2, y2, color=NAVY, width_pt=0.75):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Pt(width_pt)
    return ln


def add_picture_fit(slide, path: str | Path, x, y, w, h, align: str = "center"):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def add_section_header(slide, x, y, w, title, number: int | None = None, *, size=9.5, color=NAVY):
    """Bold header with a thin underline and an optional numbered square at the right end."""
    max_chars = int((w - (0.4 if number else 0)) * 14 / (size / 9.5))
    if len(title) > max_chars:
        size = max(7.5, size * max_chars / len(title))
    add_text(slide, x, y, w - (0.4 if number else 0), 0.26, title, size=size, bold=True, color=color, anchor=MSO_ANCHOR.BOTTOM, margin=0.02,
             wrap=False)
    add_line(slide, x, y + 0.27, x + w - (0.4 if number else 0), y + 0.27, color=NAVY, width_pt=1.0)
    if number is not None:
        sq = add_rect(slide, x + w - 0.3, y + 0.04, 0.28, 0.28, fill=NAVY)
        tf = sq.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(number)
        _style_run(r, 9, True, WHITE, FONT_BODY)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return y + 0.34


def add_colored_header(slide, x, y, w, title, fill=NAVY, h=0.27):
    bar = add_rect(slide, x, y, w, h, fill=fill)
    tf = bar.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    _style_run(r, 9.5, True, WHITE, FONT_BODY)
    return y + h


def _cell_text(cell, text, *, size=7.0, bold=False, color=BLACK, align=PP_ALIGN.RIGHT, italic=False, fill: str | None = None, font=FONT_BODY):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = "" if text is None else str(text)
    _style_run(r, size, bold, color, font, italic)
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.01)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(WHITE)


def _cell_border(cell, *, bottom: str | None = None, top: str | None = None, width_pt=0.5):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag, color in (("a:lnL", None), ("a:lnR", None), ("a:lnT", top), ("a:lnB", bottom)):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
        ln = tcPr.makeelement(qn(tag), {"w": str(int(Pt(width_pt))) if color else "0"})
        if color:
            sf = ln.makeelement(qn("a:solidFill"), {})
            sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
            ln.append(sf)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.append(ln)


def add_table(slide, x, y, w, data: list[list], *, col_widths: Sequence[float] | None = None, row_h=0.19, font_size=7.0,
              header_rows=1, header_fill=NAVY, bold_rows: Sequence[int] = (), italic_rows: Sequence[int] = (),
              first_col_align=PP_ALIGN.LEFT, cell_fills: dict | None = None, cell_colors: dict | None = None,
              band_rows: Sequence[int] = (), band_fill=LIGHT_GREY, col_fills: dict | None = None, left_cols: int = 1):
    rows, cols = len(data), max(len(r) for r in data)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for el in list(tblPr):
        if el.tag == qn("a:tableStyleId"):
            tblPr.remove(el)
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for i in range(rows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(cols):
            val = data[i][j] if j < len(data[i]) else ""
            is_header = i < header_rows
            fill = header_fill if is_header else (band_fill if i in band_rows else None)
            if col_fills and j in col_fills and not is_header:
                fill = col_fills[j]
            if cell_fills and (i, j) in cell_fills:
                fill = cell_fills[(i, j)]
            color = WHITE if is_header else BLACK
            if cell_colors and (i, j) in cell_colors:
                color = cell_colors[(i, j)]
            _cell_text(tbl.cell(i, j), val, size=font_size, bold=is_header or i in bold_rows, italic=i in italic_rows,
                       color=color, align=first_col_align if j < left_cols else PP_ALIGN.RIGHT, fill=fill)
            _cell_border(tbl.cell(i, j), bottom=MID_GREY if (i == rows - 1 or i in bold_rows) else None,
                         top=NAVY if i == header_rows and header_rows else None)
    return shape


# =============================================================================
# deck context
# =============================================================================
@dataclass
class DeckContext:
    result: object
    charts: dict[str, Path] = field(default_factory=dict)
    page: int = 0
    template_used: bool = False


def _footer(slide, ctx: DeckContext, sources: str):
    r = ctx.result
    add_line(slide, LM, FOOTER_Y, SW - 1.9, FOOTER_Y, color=NAVY, width_pt=0.75)
    add_text(slide, LM, FOOTER_Y + 0.02, SW - 2.5, 0.34, [f"Sources: {sources}", f"{r.cfg.brand.footer_note}"], size=6.5, italic=True,
             color=DARK_GREY, margin=0.0, markup=False)
    add_line(slide, SW - 1.75, FOOTER_Y, SW - LM, FOOTER_Y, color=NAVY, width_pt=0.75)
    add_text(slide, SW - 1.75, FOOTER_Y + 0.04, 1.0, 0.3, r.cfg.brand.name, size=7.5, bold=True, color=NAVY, font=FONT_TITLE, margin=0.0,
             markup=False, wrap=True)
    add_text(slide, SW - 0.72, FOOTER_Y + 0.04, 0.25, 0.3, str(ctx.page), size=8, color=NAVY, align=PP_ALIGN.RIGHT, margin=0.0, markup=False)


def _header(slide, ctx: DeckContext, tag: str, title: str, subtitle: str, headline: str | None = None) -> float:
    add_text(slide, LM, 0.18, CW, 0.24, tag, size=8, bold=True, color=NAVY, margin=0.0, markup=False)
    add_text(slide, LM, 0.42, CW, 0.55, title, size=26, color=NAVY, font=FONT_TITLE, margin=0.0, markup=False)
    add_text(slide, LM, 0.98, CW, 0.3, subtitle, size=12, color=NAVY_TEXT, margin=0.0, markup=False)
    if headline:
        bar = add_rect(slide, LM, 1.32, CW, 0.34, fill=NAVY)
        tf = bar.text_frame
        tf.margin_left = Inches(0.08)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        for t, _bold in _runs(headline):
            r = p.add_run()
            r.text = t
            _style_run(r, 9.5, True, WHITE, FONT_BODY)
        return 1.78
    return 1.40


def _new_slide(prs, ctx: DeckContext):
    layout = _blank_layout(prs)
    slide = prs.slides.add_slide(layout)
    ctx.page += 1
    return slide


def _blank_layout(prs):
    for lay in prs.slide_layouts:
        if lay.name.strip().lower() == "blank":
            return lay
    return prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]


def _clear_template_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


# =============================================================================
# slides
# =============================================================================
def _cover(prs, ctx: DeckContext):
    r = ctx.result
    slide = _new_slide(prs, ctx)
    add_rect(slide, 0, 0, SW, SH, fill=NAVY)
    add_rect(slide, 0, SH * 0.62, SW, SH * 0.38, fill="0A3A5E")
    add_text(slide, 0.95, 0.85, 8, 0.6, r.cfg.brand.name, size=26, color=WHITE, font=FONT_TITLE, margin=0, markup=False)
    add_text(slide, 0.98, 1.42, 8, 0.4, r.cfg.brand.analyst or "Equity Research", size=14, color=WHITE, margin=0, markup=False)
    add_line(slide, 1.33, 2.72, 2.45, 2.72, color=WHITE, width_pt=1.0)
    tagline = r.narrative.get("cover_tagline") or r.cfg.tagline or ""
    title = f"{r.cfg.display_short} – {tagline}" if tagline else r.name
    add_text(slide, 1.33, 2.95, 11, 1.2, title, size=34, color=WHITE, font=FONT_TITLE, margin=0, markup=False)
    sub = r.cfg.event_title or f"Initiation of coverage – {r.recommendation.rating}, target price {r.price_currency} {fmt_num(r.recommendation.target_price, 2)}"
    add_text(slide, 1.33, 4.55, 11, 0.5, sub, size=16, color=WHITE, font=FONT_TITLE, margin=0, markup=False)
    add_text(slide, 1.33, 6.2, 6, 0.3, r.cfg.date_label or r.as_of.strftime("%d.%m.%Y"), size=11, color=WHITE, margin=0, markup=False)
    add_text(slide, SW - 3.5, SH - 0.55, 3.0, 0.3, "Private and Confidential", size=7.5, color=WHITE, align=PP_ALIGN.RIGHT, margin=0, markup=False)


def _team(prs, ctx: DeckContext):
    r = ctx.result
    team = r.cfg.team
    if not team:
        return
    slide = _new_slide(prs, ctx)
    add_line(slide, LM, 0.36, LM + 0.6, 0.36, color=NAVY, width_pt=1.0)
    add_text(slide, LM, 0.45, CW, 0.55, "Case-team" if len(team) > 1 else "Analyst", size=26, color=NAVY, font=FONT_TITLE, margin=0, markup=False)
    add_text(slide, LM, 1.0, CW, 0.3, f"Presenting a model-driven investment case on {r.name}.", size=12, color=NAVY_TEXT, margin=0, markup=False)
    cols = 2
    cw, ch = 6.0, 2.6
    for i, m in enumerate(team[:4]):
        x = LM + (i % cols) * (cw + 0.4)
        y = 1.75 + (i // cols) * (ch + 0.25)
        if m.photo and Path(m.photo).exists():
            add_picture_fit(slide, m.photo, x, y, 1.55, 1.9, align="left")
        else:
            ph = add_rect(slide, x, y, 1.55, 1.9, fill=LIGHT_GREY)
            initials = "".join(w[0] for w in m.name.split()[:2]).upper()
            tf = ph.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            rr = p.add_run()
            rr.text = initials
            _style_run(rr, 28, True, GREY, FONT_TITLE)
        bar = add_rect(slide, x, y + 1.95, 3.4, 0.34, fill=LIGHT_BLUE)
        tf = bar.text_frame
        tf.margin_left = Inches(0.1)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        rr = p.add_run()
        rr.text = m.name
        _style_run(rr, 11, True, WHITE, FONT_BODY)
        lines = [f"University: {m.school}" if m.school else "", f"Year of study: {m.year}" if m.year else "",
                 f"Experience: {m.experience}" if m.experience else ""]
        add_text(slide, x + 1.75, y + 0.15, cw - 1.8, 1.7, [ln for ln in lines if ln], size=11, color=BLACK, margin=0, markup=False, line_spacing=1.4)
    _footer(slide, ctx, r.cfg.sources_note)


def _company(prs, ctx: DeckContext):
    r = ctx.result
    n = r.narrative
    info = r.snapshot.info
    slide = _new_slide(prs, ctx)
    top = _header(slide, ctx, f"{SECTION_TAGS[0]} of {r.cfg.display_short}", "Company overview", n.get("company_subtitle", ""),
                  n.get("company_headline"))
    col_w, gap = 3.95, 0.27
    xs = [LM + i * (col_w + gap) for i in range(3)]
    row1_h, row2_top = 2.05, 4.30
    # --- box 1: snapshot
    y = add_section_header(slide, xs[0], top, col_w, "Company snapshot", 1)
    add_rect(slide, xs[0], y, col_w, row1_h, fill=LIGHT_GREY)
    facts = [
        ("Sector", r.cfg.sector_label or info.get("industry", "")),
        ("Footprint", r.cfg.footprint or info.get("country", "")),
        ("Business model", r.cfg.business_model or (info.get("industry") or "")),
        ("Scale", f"Employees: {fmt_num(info.get('fullTimeEmployees'))} | Revenue {r.last_actual_year}A: {r.currency} {fmt_num(r.hist['revenue'].iloc[-1])}m"),
        ("Market cap", f"{r.currency} {fmt_num(r.market_cap)}m | EV {r.currency} {fmt_num(r.enterprise_value)}m"
                       + (f" | reports in {r.currency}, trades in {r.price_currency}" if r.dual_currency else "")),
        ("Share price", f"{fmt_price(r.price, r.price_currency)} | 52w range {fmt_num(info.get('fiftyTwoWeekLow'), 1)}–{fmt_num(info.get('fiftyTwoWeekHigh'), 1)}"),
    ]
    add_text(slide, xs[0] + 0.05, y + 0.06, col_w - 0.1, row1_h - 0.1, [f"**{k}:** {v}" for k, v in facts], size=8.5, color=BLACK, line_spacing=1.25)
    # --- box 2: highlights
    y = add_section_header(slide, xs[1], top, col_w, "Investment highlights", 2)
    add_rect(slide, xs[1], y, col_w, row1_h, fill=LIGHT_GREY)
    add_bullets(slide, xs[1] + 0.02, y + 0.04, col_w - 0.04, row1_h - 0.08, n.get("highlights", []), size=8.5)
    # --- box 3: management
    y = add_section_header(slide, xs[2], top, col_w, "Management team", 3)
    add_rect(slide, xs[2], y, col_w, row1_h, fill=LIGHT_GREY)
    officers = r.snapshot.officers[:4] or [{"name": "n.a.", "title": ""}]
    oy = y + 0.12
    for o in officers:
        circ = add_rect(slide, xs[2] + 0.15, oy, 0.36, 0.36, fill=NAVY, shape=MSO_SHAPE.OVAL)
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        rr = p.add_run()
        rr.text = "".join(w[0] for w in o["name"].replace("Mr.", "").replace("Ms.", "").split()[:2]).upper()
        _style_run(rr, 8, True, WHITE, FONT_BODY)
        clean_name = re.sub(r"^(Mr\.|Ms\.|Mrs\.|Dr\.)\s*", "", o["name"])
        add_text(slide, xs[2] + 0.6, oy - 0.02, col_w - 0.7, 0.42, [f"**{clean_name}**", _short_title(o["title"])], size=8, color=BLACK, margin=0.02, line_spacing=1.0)
        oy += 0.47
    if n.get("management_note"):
        add_text(slide, xs[2] + 0.1, y + row1_h - 0.36, col_w - 0.2, 0.34, n["management_note"], size=7.5, italic=True, color=DARK_GREY, margin=0.02)
    # --- row 2 charts
    y = add_section_header(slide, xs[0], row2_top, col_w, f"Revenue and EBITDA ({r.units_label})", 4)
    add_rect(slide, xs[0], y, col_w, 2.3, fill=LIGHT_GREY)
    add_picture_fit(slide, ctx.charts["rev_margin"], xs[0] + 0.05, y + 0.05, col_w - 0.1, 2.2)
    y = add_section_header(slide, xs[1], row2_top, col_w, f"Share price vs {r.cfg.index_name or r.cfg.index} (5 years)", 5)
    add_rect(slide, xs[1], y, col_w, 2.3, fill=LIGHT_GREY)
    if "rel_perf" in ctx.charts:
        add_picture_fit(slide, ctx.charts["rel_perf"], xs[1] + 0.05, y + 0.05, col_w - 0.1, 2.2)
    y = add_section_header(slide, xs[2], row2_top, col_w, "Key ratios", 6)
    add_rect(slide, xs[2], y, col_w, 2.3, fill=LIGHT_GREY)
    _key_ratio_table(slide, r, xs[2] + 0.05, y + 0.08, col_w - 0.1)
    _footer(slide, ctx, r.cfg.sources_note)


def _short_title(title: str) -> str:
    t = title.replace("Chief Executive Officer", "CEO").replace("Chief Financial Officer", "CFO").replace("Chief Operating Officer", "COO")
    t = t.replace("Chief Commercial Officer", "CCO").replace("Chief Technology Officer", "CTO").replace("Chief Digital Officer", "CDO")
    return t[:48]


def _key_ratio_table(slide, r, x, y, w):
    h = r.hist.tail(3)
    f = r.forecast[r.forecast.index != "TV"].head(2)
    cols = [f"{int(i)}A" for i in h.index] + [f"{int(i)}E" for i in f.index]
    rows = [["", *cols]]
    def row(label, key, fmt):
        vals = [fmt(v) for v in list(h[key]) + list(f[key] if key in f else [np.nan] * len(f))]
        rows.append([label, *vals])
    row("Revenue growth", "growth", fmt_pct)
    row("EBITDA margin", "ebitda_margin", fmt_pct)
    row("EBIT margin", "ebit_margin", fmt_pct)
    row("FCF margin", "ufcf_margin", fmt_pct)
    rows.append(["ROIC", *[fmt_pct(v) for v in h["roic"]], *["–"] * len(f)])
    rows.append(["Net debt / EBITDA", *[fmt_leverage(v) for v in h["nd_to_ebitda"]], *["–"] * len(f)])
    rows.append(["Payout ratio", *[fmt_pct(v, 0) for v in h["payout"]], *["–"] * len(f)])
    n = len(cols)
    add_table(slide, x, y, w, rows, col_widths=[w - n * 0.5] + [0.5] * n, row_h=0.24, font_size=7.5, band_rows=[2, 4, 6])


def _market(prs, ctx: DeckContext):
    r = ctx.result
    n = r.narrative
    slide = _new_slide(prs, ctx)
    top = _header(slide, ctx, SECTION_TAGS[1], "Market overview", n.get("market_subtitle", ""))
    chart_keys = [k for k in ctx.charts if k.startswith("market_")]
    col_w, gap = 4.0, 0.22
    xs = [LM, LM + col_w + gap]
    ys = [top, top + 2.75]
    titles = ctx.charts.get("_market_titles", {})
    for i, key in enumerate(chart_keys[:4]):
        x, y = xs[i % 2], ys[i // 2]
        yy = add_section_header(slide, x, y, col_w, titles.get(key, key), i + 1)
        add_rect(slide, x, yy, col_w, 2.28, fill=LIGHT_GREY)
        add_picture_fit(slide, ctx.charts[key], x + 0.05, yy + 0.05, col_w - 0.1, 2.18)
    # right column: opportunities / disruption / threats
    x = LM + 2 * col_w + 2 * gap
    w = SW - LM - x
    blocks = [("Opportunities", NAVY, n.get("opportunities", [])), ("Disruption", ACCENT_BLUE, n.get("disruption", [])),
              ("Threats", PALE_BLUE, n.get("threats", []))]
    y = top
    block_h = (FOOTER_Y - 0.1 - top - 2 * 0.15) / 3
    for title, color, items in blocks:
        yy = add_colored_header(slide, x, y, w, title, fill=color)
        add_rect(slide, x, yy, w, block_h - 0.27, fill=LIGHT_GREY)
        add_bullets(slide, x + 0.02, yy + 0.04, w - 0.04, block_h - 0.35, items, size=8.5, color=NAVY_TEXT)
        y += block_h + 0.15
    _footer(slide, ctx, r.cfg.market.sources or r.cfg.sources_note)


def _financials(prs, ctx: DeckContext):
    r = ctx.result
    n = r.narrative
    slide = _new_slide(prs, ctx)
    top = _header(slide, ctx, SECTION_TAGS[2], "Financials and estimates", n.get("financials_subtitle", ""))
    left_w = 9.35
    # --- main table
    ct = r.combined_table()
    years = list(ct.index)
    n_hist = sum(1 for y in years if y.endswith("A"))
    header1 = ["", *(["HISTORICAL"] + [""] * (n_hist - 1)), *(["FORECAST"] + [""] * (len(years) - n_hist - 1))]
    header2 = [r.units_label, *years]
    rows = [header1, header2]
    spec = [("Revenue", "revenue", fmt_num, True), ("Growth", "growth", fmt_pct, False), ("EBITDA", "ebitda", fmt_num, True),
            ("Margin", "ebitda_margin", fmt_pct, False), ("EBIT", "ebit", fmt_num, True), ("Margin", "ebit_margin", fmt_pct, False),
            ("NOPAT", "nopat", fmt_num, True), ("Margin", "nopat_margin", fmt_pct, False), ("Unlevered FCF", "ufcf", fmt_num, True),
            ("FCF margin", "ufcf_margin", fmt_pct, False)]
    bold_rows, italic_rows = [], []
    for label, key, fmt, bold in spec:
        rows.append([label, *[fmt(v) for v in ct[key]]])
        (bold_rows if bold else italic_rows).append(len(rows) - 1)
    ncol = len(years)
    cw = (left_w - 1.15) / ncol
    fills = {(0, j): NAVY for j in range(ncol + 1)}
    for j in range(1, n_hist + 1):
        fills[(0, j)] = GREY
    col_fills = {j: "F4F6F8" for j in range(n_hist + 1, ncol + 1)}
    shape = add_table(slide, LM, top, left_w, rows, col_widths=[1.15] + [cw] * ncol, row_h=0.185, font_size=7, header_rows=2,
                      bold_rows=bold_rows, italic_rows=italic_rows, cell_fills=fills, col_fills=col_fills)
    tbl = shape.table
    if n_hist > 1:
        tbl.cell(0, 1).merge(tbl.cell(0, n_hist))
    if ncol - n_hist > 1:
        tbl.cell(0, n_hist + 1).merge(tbl.cell(0, ncol))
    for j in (1, n_hist + 1):
        for p in tbl.cell(0, j).text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    y = top + 0.185 * len(rows) + 0.15
    # --- multiples table
    fm = r.forward_multiples
    mrows = [["Multiples (at current price)", *list(fm.columns)]]
    for label in fm.index:
        f = (lambda v: fmt_pct(v)) if "yield" in label.lower() else fmt_mult
        mrows.append([label, *[f(v) for v in fm.loc[label]]])
    add_table(slide, LM, y, 1.15 + 0.85 * len(fm.columns), mrows, col_widths=[1.15] + [0.85] * len(fm.columns), row_h=0.185, font_size=7)
    # consensus box next to multiples table
    cx = LM + 1.15 + 0.85 * len(fm.columns) + 0.25
    cw2 = left_w - (cx - LM)
    pt = r.snapshot.price_targets or {}
    rs = r.snapshot.recommendations
    cons_lines = [f"**Consensus** ({fmt_num(r.snapshot.info.get('numberOfAnalystOpinions'))} analysts)"]
    if pt.get("mean"):
        cons_lines.append(f"Target: {r.price_currency} {fmt_num(pt.get('mean'), 2)} (range {fmt_num(pt.get('low'), 0)}–{fmt_num(pt.get('high'), 0)})")
    if rs is not None and not rs.empty:
        row0 = rs.iloc[0]
        cons_lines.append(f"Ratings: {int(row0.get('strongBuy', 0) + row0.get('buy', 0))} buy / {int(row0.get('hold', 0))} hold / "
                          f"{int(row0.get('sell', 0) + row0.get('strongSell', 0))} sell")
    cons_lines.append(f"**Engine:** {r.recommendation.rating}, TP {r.price_currency} {fmt_num(r.recommendation.target_price, 2)} "
                      f"({fmt_pct(r.recommendation.upside, 0, sign=True)})")
    add_rect(slide, cx, y, cw2, 0.185 * len(mrows), fill=LIGHT_GREY)
    add_text(slide, cx + 0.05, y + 0.02, cw2 - 0.1, 0.185 * len(mrows), cons_lines, size=7.5, color=BLACK, line_spacing=1.1)
    y += 0.185 * len(mrows) + 0.2
    # --- multiples panels (or the forecast drivers when there is no peer group)
    panel_w = (left_w - 0.3) / 3
    ph = FOOTER_Y - 0.12 - y
    if any(f"mult_{k}" in ctx.charts for k in ("ev_ebitda", "ev_ebit", "pe")):
        for i, key in enumerate(("ev_ebitda", "ev_ebit", "pe")):
            ck = f"mult_{key}"
            if ck in ctx.charts:
                add_picture_fit(slide, ctx.charts[ck], LM + i * (panel_w + 0.15), y, panel_w, ph)
    else:
        d = r.drivers
        yrs = [f"{yy}E" for yy in d.years]
        drows = [["Forecast drivers", *yrs],
                 ["Revenue growth", *[fmt_pct(v) for v in d.revenue_growth]],
                 ["EBITDA margin", *[fmt_pct(v) for v in d.ebitda_margin]],
                 ["D&A % of revenue", *[fmt_pct(d.da_pct)] * len(yrs)],
                 ["Capex % of revenue", *[fmt_pct(d.capex_at(j)) for j in range(len(yrs))]],
                 ["Net working capital % of revenue", *[fmt_pct(d.nwc_pct)] * len(yrs)],
                 ["Tax rate", *[fmt_pct(d.tax_rate)] * len(yrs)]]
        yy2 = add_section_header(slide, LM, y, left_w, "Forecast drivers (edit them in the case config)")
        add_table(slide, LM, yy2 + 0.05, min(left_w, 2.2 + 0.85 * len(yrs)), drows, col_widths=[2.2] + [0.85] * len(yrs), row_h=0.18, font_size=7.5,
                  band_rows=[2, 4, 6])
        notes = " ".join(d.notes)
        room = FOOTER_Y - 0.06 - (yy2 + 0.05 + 0.18 * len(drows) + 0.04)
        if room >= 0.2 and notes:
            max_chars = int(room / 0.13) * 150
            if len(notes) > max_chars:
                notes = notes[: max_chars - 1].rsplit(" ", 1)[0] + "…"
            add_text(slide, LM, yy2 + 0.05 + 0.18 * len(drows) + 0.04, left_w, room, notes, size=7, italic=True, color=DARK_GREY, markup=False)
    # --- commentary
    x = LM + left_w + 0.25
    w = SW - LM - x
    add_text(slide, x, top - 0.05, w, 0.3, "Financial commentary", size=10, bold=True, color=NAVY, margin=0.02, markup=False)
    add_line(slide, x, top + 0.25, x + w, top + 0.25, color=NAVY, width_pt=1.0)
    add_bullets(slide, x, top + 0.3, w, FOOTER_Y - top - 0.4, n.get("financial_commentary", []), size=8.5, space_after=5)
    _footer(slide, ctx, r.cfg.sources_note + "; peer multiples on last reported fiscal year")


def _valuation(prs, ctx: DeckContext):
    r = ctx.result
    n = r.narrative
    rec = r.recommendation
    slide = _new_slide(prs, ctx)
    top = _header(slide, ctx, SECTION_TAGS[3], "Valuation and recommendation", n.get("valuation_subtitle", ""),
                  n.get("valuation_headline") or None)
    # ---------------- left column
    lx, lw = LM, 4.25
    tp_label = "12m TP" if rec.horizon_months else "TP"
    agree = getattr(r.crosscheck, "agreement", "n/a") if r.crosscheck is not None else "n/a"
    agree_txt = f" · {agree} agreement" if agree in {"low", "medium"} else ""
    y = add_section_header(slide, lx, top, lw, f"{rec.rating} – {tp_label} {r.price_currency} {fmt_num(rec.target_price, 2)}{agree_txt}")
    add_rect(slide, lx, y, lw, 1.85, fill=LIGHT_GREY, line=DARK_GREY, line_w=0.75)
    add_bullets(slide, lx + 0.02, y + 0.04, lw - 0.04, 1.8, n.get("thesis", []), size=8)
    y2 = y + 1.85 + 0.22
    y3 = add_section_header(slide, lx, y2, lw, "Valuation summary – Football field")
    add_rect(slide, lx, y3, lw, FOOTER_Y - 0.12 - y3, fill=WHITE, line=NAVY, line_w=0.75)
    if "football" in ctx.charts:
        add_picture_fit(slide, ctx.charts["football"], lx + 0.05, y3 + 0.05, lw - 0.1, FOOTER_Y - 0.22 - y3)
    # ---------------- middle column: sensitivity
    mx, mw = LM + lw + 0.3, 5.75
    y = add_section_header(slide, mx, top, mw, "DCF valuation – Sensitivity analysis (value per share)")
    s = r.dcf.sensitivity
    k = len(s) // 2
    header = ["Terminal growth \\ WACC", *[fmt_pct(w, 2) for w in s.columns]]
    rows = [header]
    fills, colors = {}, {}
    for i, g in enumerate(s.index):
        rows.append([fmt_pct(g, 2), *[fmt_num(v, 1) for v in s.loc[g]]])
        fills[(i + 1, 0)] = NAVY
        colors[(i + 1, 0)] = WHITE
        for j in range(len(s.columns)):
            if abs(i - k) <= 1 and abs(j - k) <= 1:
                fills[(i + 1, j + 1)] = "DCE6F0"
            if i == k and j == k:
                fills[(i + 1, j + 1)] = LIGHT_BLUE
    ncol = len(s.columns)
    add_table(slide, mx, y, mw, rows, col_widths=[1.35] + [(mw - 1.35) / ncol] * ncol, row_h=0.235, font_size=7.5,
              cell_fills=fills, cell_colors=colors, bold_rows=[k + 1])
    y = y + 0.235 * len(rows) + 0.2
    implied = n.get("market_implied", [])
    yy = add_section_header(slide, mx, y, mw, "Multiple valuation and what the price implies" if implied else "Multiple valuation")
    add_rect(slide, mx, yy, mw, FOOTER_Y - 0.12 - yy, fill=LIGHT_GREY, line=DARK_GREY, line_w=0.75)
    add_bullets(slide, mx + 0.02, yy + 0.04, mw - 0.04, FOOTER_Y - 0.2 - yy, [*n.get("multiples_commentary", []), *implied], size=8,
                space_after=3)
    # ---------------- right column: scenarios and multiple-implied values
    rx = mx + mw + 0.25
    rw = SW - LM - rx
    add_rect(slide, rx, top, rw, 2.45, fill=WHITE, line=NAVY, line_w=0.75)
    top_chart = "scenarios" if "scenarios" in ctx.charts else "bridge_dcf"
    if top_chart in ctx.charts:
        add_picture_fit(slide, ctx.charts[top_chart], rx + 0.05, top + 0.05, rw - 0.1, 2.35)
    yb = top + 2.45 + 0.25
    if "bridge_mult" in ctx.charts:
        add_rect(slide, rx, yb, rw, FOOTER_Y - 0.12 - yb, fill=WHITE, line=NAVY, line_w=0.75)
        add_picture_fit(slide, ctx.charts["bridge_mult"], rx + 0.05, yb + 0.05, rw - 0.1, FOOTER_Y - 0.22 - yb)
    _footer(slide, ctx, f"{r.cfg.sources_note}; WACC {fmt_pct(r.wacc.wacc)} ({r.wacc.beta_source}); terminal growth {fmt_pct(r.drivers.terminal_growth)}")


# =============================================================================
# appendix slides
# =============================================================================
def _appendix_header(slide, ctx: DeckContext, number: str, title: str, subtitle: str) -> float:
    return _header(slide, ctx, f"Appendix #{number}", title, subtitle)


def _cap_first(text: str) -> str:
    """Upper-case the first letter only (str.capitalize would lower-case 'USD')."""
    return text[:1].upper() + text[1:]


def _appendix_dcf(prs, ctx: DeckContext):
    """WACC build-up, DCF build-up and the 12-month target price mechanics."""
    r = ctx.result
    w, d, rec = r.wacc, r.dcf, r.recommendation
    pccy = r.price_currency
    slide = _new_slide(prs, ctx)
    top = _appendix_header(slide, ctx, "4.1", "Cost of capital and DCF build-up",
                           f"WACC {fmt_pct(w.wacc, 2)} and terminal growth {fmt_pct(d.terminal_growth, 2)} give a fair value of "
                           f"{pccy} {fmt_num(d.value_per_share, 2)} per share")
    # ---- WACC
    lx, lw = LM, 3.7
    y = add_section_header(slide, lx, top, lw, "WACC build-up", 1)
    rows = [["Component", "Value"],
            ["Risk-free rate", fmt_pct(w.risk_free, 2)], ["Beta (raw)", fmt_num(w.beta_raw, 2)], ["Beta (adjusted, used)", fmt_num(w.beta_used, 2)],
            ["Equity risk premium", fmt_pct(w.equity_risk_premium, 2)], ["Size premium", fmt_pct(w.size_premium, 2)],
            ["Cost of equity", fmt_pct(w.cost_of_equity, 2)], ["Cost of debt (pre-tax)", fmt_pct(w.cost_of_debt_pretax, 2)],
            ["Tax rate", fmt_pct(w.tax_rate, 1)], ["Cost of debt (after tax)", fmt_pct(w.cost_of_debt_after_tax, 2)],
            ["Weight of equity", fmt_pct(w.weight_equity, 1)], ["Weight of debt", fmt_pct(w.weight_debt, 1)], ["WACC", fmt_pct(w.wacc, 2)]]
    add_table(slide, lx, y + 0.03, lw, rows, col_widths=[2.5, 1.2], row_h=0.235, font_size=8, bold_rows=[6, 12], band_rows=[2, 4, 8, 10])
    notes = [f"Beta: {w.beta_source}."]
    if w.peer_beta is not None:
        pb = w.peer_beta
        notes.append(f"Cross-check: bottom-up beta {pb.relevered_adjusted:.2f} ({pb.n} peers, unlevered median {pb.unlevered_median:.2f}).")
    notes += [f"{_cap_first(w.size_premium_source or 'size premium from config')}.", f"Cost of debt: {w.cost_of_debt_source}."]
    add_text(slide, lx, y + 0.03 + 0.235 * len(rows) + 0.08, lw, 1.2, notes, size=7.5, italic=True, color=DARK_GREY, markup=False, line_spacing=1.1)
    # ---- DCF
    mx = lx + lw + 0.3
    mw = 5.6
    y = add_section_header(slide, mx, top, mw, f"DCF build-up ({r.units_label})", 2)
    yrs = [f"{yy}E" for yy in d.years]
    n = len(yrs)
    cw = (mw - 1.55) / n
    rows = [["", *yrs], ["Unlevered FCF", *[fmt_num(v) for v in d.fcf]]]
    stub = bool(d.fcf_weights) and d.fcf_weights[0] < 0.999
    if stub:
        rows.append(["Share still to come", *[fmt_pct(v, 0) for v in d.fcf_weights]])
    rows += [["Years to discount", *[f"{v:.2f}" for v in d.discount_periods]] if d.discount_periods else ["Discount period", *[""] * n],
             ["Discount factor", *[f"{v:.3f}" for v in d.discount_factors]], ["PV of FCF", *[fmt_num(v) for v in d.pv_fcf]]]
    add_table(slide, mx, y + 0.03, mw, rows, col_widths=[1.55] + [cw] * n, row_h=0.235, font_size=8, bold_rows=[len(rows) - 1])
    yb = y + 0.03 + 0.235 * len(rows) + 0.15
    tv_label = "Terminal value (value driver)" if d.terminal_method == "value_driver" else "Terminal value (Gordon growth)"
    nd_date = f" ({pd.Timestamp(r.latest.net_debt_date):%d.%m.%Y})" if (r.latest is not None and r.latest.net_debt_date) else ""
    bridge = [["Bridge to value per share", r.units_label],
              ["Sum of PV of FCF", fmt_num(d.sum_pv_fcf)], [tv_label, fmt_num(d.terminal_value)],
              ["PV of terminal value", fmt_num(d.pv_terminal)], ["Enterprise value", fmt_num(d.enterprise_value)],
              [f"Less: net debt{nd_date}", fmt_num(-d.net_debt)], ["Less: minority interest", fmt_num(-d.minorities)],
              ["Equity value", fmt_num(d.equity_value)], ["Shares outstanding (m)", fmt_num(r.shares_real, 1)]]
    if r.dual_currency:
        bridge.append([f"FX ({r.currency} per {pccy})", f"{r.fx_reporting_per_listing:.4f}"])
    vd = f", {r.valuation_date:%d.%m.%Y}" if r.valuation_date else ""
    bridge.append([f"Fair value per share ({pccy}{vd})", fmt_num(d.value_per_share, 2)])
    bold = [4, 7, len(bridge) - 1]
    if rec.horizon_months:
        bridge += [[f"x (1 + cost of equity {fmt_pct(rec.cost_of_equity, 2)})", fmt_num(rec.fair_value * (1 + rec.cost_of_equity) ** (rec.horizon_months / 12), 2)],
                   ["Less: expected dividend per share", fmt_num(-rec.dps, 2)],
                   [f"{rec.horizon_months}-month target price ({pccy}, rounded)", fmt_num(rec.target_price, 2)]]
        bold.append(len(bridge) - 1)
    add_table(slide, mx, yb, mw, bridge, col_widths=[mw - 1.3, 1.3], row_h=0.22, font_size=8, bold_rows=bold)
    # ---- reverse DCF
    rx = mx + mw + 0.3
    rw = SW - LM - rx
    y = add_section_header(slide, rx, top, rw, "What the share price implies", 3)
    rv = r.reverse_dcf
    box_h = 2.35
    add_rect(slide, rx, y, rw, box_h, fill=LIGHT_GREY)
    items = []
    if rv is not None:
        items.append(f"**Implied WACC:** {fmt_pct(rv.implied_wacc, 2) if rv.implied_wacc is not None else 'n.a.'} vs our {fmt_pct(rv.base_wacc, 2)}")
        items.append(f"**Implied terminal growth:** {fmt_pct(rv.implied_terminal_growth, 2) if rv.implied_terminal_growth is not None else 'outside −10% to WACC'} "
                     f"vs our {fmt_pct(rv.base_terminal_growth, 2)}")
        items.append(f"**Implied {d.years[-1]}E EBITDA margin:** {fmt_pct(rv.implied_final_margin) if rv.implied_final_margin is not None else 'n.a.'} "
                     f"vs our {fmt_pct(rv.base_final_margin)}")
    items.append("Each figure solves for the one input that makes the DCF equal today's price, holding the others at the base case")
    add_bullets(slide, rx + 0.02, y + 0.04, rw - 0.04, box_h - 0.08, items, size=8)
    y2 = y + box_h + 0.2
    y2 = add_section_header(slide, rx, y2, rw, "Key checks", 4)
    add_rect(slide, rx, y2, rw, FOOTER_Y - 0.12 - y2, fill=LIGHT_GREY)
    add_bullets(slide, rx + 0.02, y2 + 0.04, rw - 0.04, FOOTER_Y - 0.2 - y2, [
        f"**Terminal value** is {fmt_pct(d.tv_share_of_ev, 0)} of enterprise value",
        f"**Implied exit multiple:** {fmt_mult(d.implied_exit_ev_ebitda)} EV/EBITDA vs {fmt_mult(r.comps.company.get('ev_ebitda'))} today",
        f"**WACC − g spread:** {fmt_pct(d.wacc - d.terminal_growth, 2)}",
        f"**Net debt / EBITDA:** {fmt_leverage(r.hist['nd_to_ebitda'].iloc[-1])}",
    ], size=8)
    _footer(slide, ctx, r.cfg.sources_note)


def _appendix_scenarios(prs, ctx: DeckContext):
    r = ctx.result
    if len(r.scenarios) < 2:
        return
    pccy = r.price_currency
    slide = _new_slide(prs, ctx)
    wv = r.scenario_weighted_value
    top = _appendix_header(slide, ctx, "4.2", "Scenario analysis",
                           f"Probability-weighted value of {pccy} {fmt_num(wv, 2)} per share versus a share price of {pccy} {fmt_num(r.price, 2)}"
                           if wv is not None else "Bear, base and bull cases")
    lw = 7.7
    y = add_section_header(slide, LM, top, lw, "Scenario table", 1)
    rows = [["Scenario", "Probability", "Growth shift", "Margin shift", "WACC", "Terminal g", "Revenue CAGR", f"{r.dcf.years[-1]}E margin",
             f"Value ({pccy})", "vs price"]]
    for s in r.scenarios:
        rows.append([s.name, fmt_pct(s.probability, 0), fmt_pct(s.growth_shift, 1, sign=True), fmt_pct(s.margin_shift, 1, sign=True), fmt_pct(s.wacc, 2),
                     fmt_pct(s.terminal_growth, 2), fmt_pct(s.revenue_cagr), fmt_pct(s.final_margin), fmt_num(s.value_per_share, 2),
                     fmt_pct(s.upside, 0, sign=True)])
    if wv is not None:
        rows.append(["Probability-weighted", "100%", "", "", "", "", "", "", fmt_num(wv, 2), fmt_pct(wv / r.price - 1, 0, sign=True)])
    ncol = len(rows[0])
    add_table(slide, LM, y + 0.03, lw, rows, col_widths=[1.55] + [(lw - 1.55) / (ncol - 1)] * (ncol - 1), row_h=0.27, font_size=8,
              bold_rows=[len(rows) - 1] if wv is not None else [], band_rows=[2])
    yb = y + 0.03 + 0.27 * len(rows) + 0.25
    yb = add_section_header(slide, LM, yb, lw, "What drives each case", 2)
    add_rect(slide, LM, yb, lw, FOOTER_Y - 0.12 - yb, fill=LIGHT_GREY)
    notes = list(r.narrative.get("scenario_commentary", []))
    notes.append("Shifts are additive to the base case: growth applies to every forecast year, the margin shift to the whole EBITDA margin path; "
                 "capex and working capital follow revenue as in the base case")
    add_bullets(slide, LM + 0.02, yb + 0.05, lw - 0.04, FOOTER_Y - 0.22 - yb, notes, size=8.5)
    rx = LM + lw + 0.3
    rw = SW - LM - rx
    y = add_section_header(slide, rx, top, rw, "Value per share by scenario", 3)
    add_rect(slide, rx, y, rw, FOOTER_Y - 0.12 - y, fill=WHITE, line=NAVY, line_w=0.75)
    if "scenarios_big" in ctx.charts:
        add_picture_fit(slide, ctx.charts["scenarios_big"], rx + 0.08, y + 0.08, rw - 0.16, FOOTER_Y - 0.28 - y)
    _footer(slide, ctx, r.cfg.sources_note)


def _appendix_peers(prs, ctx: DeckContext):
    r = ctx.result
    tbl = r.comps.table
    if tbl.empty:
        return
    slide = _new_slide(prs, ctx)
    top = _appendix_header(slide, ctx, "4.3", "Peer group detail",
                           "Last-twelve-month multiples where the quarters are reported; EV on today's market cap, the latest net debt and "
                           "the same lease basis as the target company")
    header = ["Company", "Group", f"Mkt cap ({r.units_label})", f"EV ({r.units_label})", "Rev. growth", "EBITDA margin", "EV/Sales", "EV/EBITDA",
              "EV/EBIT", "P/E", "P/E (NTM)"]
    rows = [header]
    for _, p in tbl.iterrows():
        rows.append([p["name"], p["group"], fmt_num(p["market_cap"]), fmt_num(p["ev"]), fmt_pct(p["revenue_growth"]), fmt_pct(p["ebitda_margin"]),
                     fmt_mult(p["ev_sales"]), fmt_mult(p["ev_ebitda"]), fmt_mult(p["ev_ebit"]), fmt_mult(p["pe"]), fmt_mult(p["fwd_pe"])])
    stats = r.comps.stats
    first_stat = len(rows)
    for label, key in (("Peer median", "All|median"), ("Peer average", "All|mean"), ("25th percentile", "All|p25"), ("75th percentile", "All|p75")):
        if key in stats.index:
            s = stats.loc[key]
            rows.append([label, "", "", "", fmt_pct(s["revenue_growth"]), fmt_pct(s["ebitda_margin"]), fmt_mult(s["ev_sales"]), fmt_mult(s["ev_ebitda"]),
                         fmt_mult(s["ev_ebit"]), fmt_mult(s["pe"]), fmt_mult(s["fwd_pe"])])
    c = r.comps.company
    h = r.hist.iloc[-1]
    L = r.latest
    own_margin = (L.ebitda / L.revenue) if (L is not None and L.revenue and L.ebitda is not None) else h["ebitda_margin"]
    rows.append([r.cfg.display_short, "Target", fmt_num(r.market_cap), fmt_num(r.enterprise_value), fmt_pct(h["growth"]), fmt_pct(own_margin),
                 fmt_mult(c.get("ev_sales")), fmt_mult(c.get("ev_ebitda")), fmt_mult(c.get("ev_ebit")), fmt_mult(c.get("pe")), fmt_mult(c.get("fwd_pe"))])
    row_h = min(0.26, (FOOTER_Y - 0.2 - top) / len(rows))
    fills = {(len(rows) - 1, j): "DCE6F0" for j in range(len(header))}
    add_table(slide, LM, top, CW, rows, col_widths=[2.4, 1.9] + [(CW - 4.3) / 9] * 9, row_h=row_h, font_size=8 if row_h >= 0.22 else 7,
              bold_rows=[first_stat, len(rows) - 1], cell_fills=fills, band_rows=list(range(2, first_stat, 2)), left_cols=2)
    bases = tbl["basis"].dropna().value_counts() if "basis" in tbl else pd.Series(dtype=int)
    basis_txt = ", ".join(f"{b} ({n})" for b, n in bases.items())
    own_basis = c.get("basis") or f"FY{r.last_actual_year}"
    _footer(slide, ctx, r.cfg.sources_note + f"; income basis: {r.cfg.display_short} {own_basis}; peers {basis_txt}; "
                                             "n.m. = not meaningful (negative or above 60x)")


def _appendix_history(prs, ctx: DeckContext):
    """The company's own multiples through time, plus the cross-check of the valuation anchors."""
    r = ctx.result
    mh = r.multiple_history
    if mh is None or not any(k in ctx.charts for k in ("hist_ev_ebitda", "hist_pe")):
        return
    pccy = r.price_currency
    slide = _new_slide(prs, ctx)
    st = mh.stats.get("ev_ebitda") or mh.stats.get("pe")
    key_label = "EV/EBITDA" if "ev_ebitda" in mh.stats else "P/E"
    top = _appendix_header(slide, ctx, "4.4", "Valuation through time",
                           f"{r.cfg.display_short} trades at {fmt_mult(st['current'])} trailing {key_label} against a {st['years']:.1f}-year median of "
                           f"{fmt_mult(st['median'])}")
    half = (CW - 0.3) / 2
    chart_h = 3.05
    for i, (key, title) in enumerate((("hist_ev_ebitda", "Trailing EV/EBITDA"), ("hist_pe", "Trailing P/E"))):
        x = LM + i * (half + 0.3)
        y = add_section_header(slide, x, top, half, title, i + 1)
        add_rect(slide, x, y, half, chart_h, fill=WHITE, line=NAVY, line_w=0.75)
        if key in ctx.charts:
            add_picture_fit(slide, ctx.charts[key], x + 0.08, y + 0.06, half - 0.16, chart_h - 0.12)
        else:
            add_text(slide, x + 0.2, y + 1.3, half - 0.4, 0.4, "Not meaningful: earnings were negative or the multiple was outside 0–80x",
                     size=9, italic=True, color=DARK_GREY, align=PP_ALIGN.CENTER, markup=False)
    yb = top + 0.34 + chart_h + 0.2
    yb = add_section_header(slide, LM, yb, CW, "Reading the bands", 3)
    add_rect(slide, LM, yb, CW, FOOTER_Y - 0.12 - yb, fill=LIGHT_GREY)
    items = []
    for key, label in (("ev_ebitda", "EV/EBITDA"), ("pe", "P/E")):
        s_ = mh.stats.get(key)
        if not s_:
            continue
        txt = (f"**{label}:** {fmt_mult(s_['current'])} today vs a median of {fmt_mult(s_['median'])} and an interquartile band of "
               f"{fmt_mult(s_['p25'])}–{fmt_mult(s_['p75'])}; {fmt_pct(s_['percentile'], 0)} of weeks were cheaper")
        iv = mh.implied.get(key)
        if iv:
            txt += f". At its own median the share is worth {pccy} {fmt_num(iv['per_share'], 2)} ({fmt_pct(iv['per_share'] / r.price - 1, 0, sign=True)} vs price)"
        items.append(txt)
    if r.crosscheck is not None and r.crosscheck.agreement != "n/a":
        items.append(f"**Cross-check – {r.crosscheck.agreement} agreement between the anchors:** {r.crosscheck.message}")
    items.append("Each weekly close is paired with the latest fiscal year that was public at the time (75-day reporting lag), on the same lease basis "
                 "as the model; net debt and the share count step once a year")
    add_bullets(slide, LM + 0.02, yb + 0.05, CW - 0.04, FOOTER_Y - 0.22 - yb, items, size=8.5)
    _footer(slide, ctx, r.cfg.sources_note)


def _appendix_multiples(prs, ctx: DeckContext):
    """Multiples in depth: the regression of EV/EBITDA on growth and margin, and forward multiples calendarised to NTM."""
    r = ctx.result
    tbl = r.comps.table
    if tbl.empty:
        return
    c, reg, stats = r.comps.company, r.comps.regression, r.comps.stats
    med = stats.loc["All|median"] if "All|median" in stats.index else None
    slide = _new_slide(prs, ctx)
    if reg is not None:
        sub = (f"Growth and margin explain {fmt_pct(reg.r2, 0)} of the dispersion in peer EV/EBITDA; fundamentals justify {fmt_mult(reg.fitted_target)} "
               f"for {r.cfg.display_short} against {fmt_mult(reg.actual_target)} today")
    else:
        sub = "Forward multiples on consensus calendarised to the next twelve months"
    top = _appendix_header(slide, ctx, "4.6", "Multiples in depth", sub)
    # ---- left: scatter + reading
    lw = 5.3
    y = add_section_header(slide, LM, top, lw, "EV/EBITDA against expected revenue growth", 1)
    chart_h = 2.85
    add_rect(slide, LM, y, lw, chart_h, fill=WHITE, line=NAVY, line_w=0.75)
    if "mult_regression" in ctx.charts:
        add_picture_fit(slide, ctx.charts["mult_regression"], LM + 0.08, y + 0.06, lw - 0.16, chart_h - 0.12)
    else:
        add_text(slide, LM + 0.2, y + 1.2, lw - 0.4, 0.5, "Too few peers with a meaningful EV/EBITDA and growth estimate for a regression",
                 size=9, italic=True, color=DARK_GREY, align=PP_ALIGN.CENTER, markup=False)
    yb = y + chart_h + 0.18
    yb = add_section_header(slide, LM, yb, lw, "Reading the multiples", 3)
    add_rect(slide, LM, yb, lw, FOOTER_Y - 0.12 - yb, fill=LIGHT_GREY)
    add_bullets(slide, LM + 0.02, yb + 0.05, lw - 0.04, FOOTER_Y - 0.22 - yb, r.narrative.get("multiples_depth", []) or ["No peer consensus available"], size=7.8)
    # ---- right: forward multiples table
    rx = LM + lw + 0.3
    rw = SW - LM - rx
    y = add_section_header(slide, rx, top, rw, "Forward multiples and quality – consensus, calendarised to the next twelve months", 2)
    header = ["Company", "EV/Sales NTM", "EV/EBITDA NTM*", "P/E FY0", "P/E FY1", "P/E NTM", "EPS growth", "PEG", "FCF yield", "Div. yield"]
    rows = [header]

    def line(name, src) -> list:
        g = src.get if isinstance(src, dict) else (lambda k, d=None: src[k] if k in src.index else d)
        return [name, fmt_mult(g("ev_sales_ntm")), fmt_mult(g("ev_ebitda_ntm")), fmt_mult(g("pe_fy0")), fmt_mult(g("pe_fy1")), fmt_mult(g("fwd_pe")),
                fmt_pct(g("eps_growth"), 0, sign=True), fmt_num(g("peg"), 2), fmt_pct(g("fcf_yield")), fmt_pct(g("div_yield"))]

    for _, p in tbl.iterrows():
        rows.append(line(p["name"], p.to_dict()))
    first_stat = len(rows)
    if med is not None:
        rows.append(line("Peer median", med))
    rows.append(line(r.cfg.display_short, c))
    row_h = min(0.26, (FOOTER_Y - 0.55 - y) / len(rows))
    fills = {(len(rows) - 1, j): "DCE6F0" for j in range(len(header))}
    # the EV/EBITDA column needs room for the word; the rest share what is left
    other = (rw - 1.45 - 0.86) / 8
    add_table(slide, rx, y + 0.03, rw, rows, col_widths=[1.45, other, 0.86] + [other] * 7, row_h=row_h, font_size=7.5 if row_h >= 0.22 else 7,
              bold_rows=[first_stat, len(rows) - 1], cell_fills=fills, band_rows=list(range(2, first_stat, 2)), left_cols=1)
    yt = y + 0.03 + row_h * len(rows) + 0.06
    add_text(slide, rx, yt, rw, FOOTER_Y - 0.12 - yt,
             [f"* NTM revenue at the LTM EBITDA margin: Yahoo Finance carries no EBITDA consensus. FY0 = fiscal year in progress; {c.get('forward_basis', '')}. "
              "EV in each company's reporting currency (market cap converted from the listing currency) with the latest net debt; FCF yield is levered, "
              "after lease payments; n.m. = negative or above the cap."], size=6.8, italic=True, color=DARK_GREY, markup=False, line_spacing=1.1)
    _footer(slide, ctx, r.cfg.sources_note + "; consensus: Yahoo Finance mean estimates; regression: OLS with an intercept over the peers shown")


def _appendix_consensus(prs, ctx: DeckContext):
    """Where we differ from consensus, where consensus is moving, what is coming up, and what moves the value."""
    r = ctx.result
    cv = r.consensus
    if (cv is None or not cv.estimates) and "tornado" not in ctx.charts:
        return
    pccy, ccy = r.price_currency, r.currency
    slide = _new_slide(prs, ctx)
    eps = [e for e in (cv.estimates if cv else []) if e.metric == "EPS" and e.diff is not None]
    if eps:
        e0 = eps[0]
        sub = (f"Our {e0.year}E EPS of {pccy} {fmt_num(e0.ours, 2)} is {fmt_pct(abs(e0.diff), 0)} {'above' if e0.diff >= 0 else 'below'} consensus; "
               f"estimate momentum is {cv.momentum}")
    else:
        sub = "What moves the value, one driver at a time"
    top = _appendix_header(slide, ctx, "4.5", "Estimates vs consensus and value drivers", sub)
    lw = 7.2
    y = add_section_header(slide, LM, top, lw, "Our estimates vs consensus", 1)
    if cv is not None and cv.estimates:
        rows = [["Estimate", "Ours", "Consensus", "Difference", "Low", "High", "Analysts"]]
        for e in sorted(cv.estimates, key=lambda e: (e.metric != "Revenue", e.year)):
            d = 2 if e.metric == "EPS" else 0
            unit = pccy if e.metric == "EPS" else f"{ccy}m"
            rows.append([f"{e.metric} {e.year}E ({unit})", fmt_num(e.ours, d), fmt_num(e.consensus, d), fmt_pct(e.diff, 1, sign=True),
                         fmt_num(e.low, d), fmt_num(e.high, d), str(e.n_analysts or "–")])
        fills = {(i, 3): ("E3EEE9" if (e.diff or 0) >= 0 else "F6E1DE") for i, e in
                 enumerate(sorted(cv.estimates, key=lambda e: (e.metric != "Revenue", e.year)), start=1) if e.diff is not None and abs(e.diff) >= 0.05}
        add_table(slide, LM, y + 0.03, lw, rows, col_widths=[2.0] + [(lw - 2.0) / 6] * 6, row_h=0.25, font_size=8, cell_fills=fills,
                  band_rows=[2, 4], left_cols=1)
        yb = y + 0.03 + 0.25 * len(rows) + 0.22
    else:
        add_text(slide, LM, y + 0.1, lw, 0.4, (cv.note if cv else "") or "No consensus estimates available on Yahoo Finance.", size=8.5, italic=True,
                 color=DARK_GREY, markup=False)
        yb = y + 0.7
    yb = add_section_header(slide, LM, yb, lw, "Variant view, estimate momentum and catalysts", 2)
    add_rect(slide, LM, yb, lw, FOOTER_Y - 0.12 - yb, fill=LIGHT_GREY)
    items = list(r.narrative.get("consensus", [])) + list(r.narrative.get("catalysts", []))
    add_bullets(slide, LM + 0.02, yb + 0.05, lw - 0.04, FOOTER_Y - 0.22 - yb, items or ["No consensus data"], size=8.2)
    rx = LM + lw + 0.3
    rw = SW - LM - rx
    y = add_section_header(slide, rx, top, rw, "What moves the value – one driver at a time", 3)
    add_rect(slide, rx, y, rw, FOOTER_Y - 0.12 - y, fill=WHITE, line=NAVY, line_w=0.75)
    if "tornado" in ctx.charts:
        add_picture_fit(slide, ctx.charts["tornado"], rx + 0.08, y + 0.08, rw - 0.16, FOOTER_Y - 0.3 - y)
    _footer(slide, ctx, r.cfg.sources_note + "; consensus: Yahoo Finance (mean of analyst estimates), revisions over 90 days; "
                                             "value drivers move one input at a time with everything else at the base case")


# =============================================================================
# charts for the deck
# =============================================================================
def render_deck_charts(result, charts_dir: Path) -> dict[str, Path]:
    r = result
    charts_dir = Path(charts_dir)
    charts_dir.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    ct = r.combined_table()
    ct = ct[ct.index != "TV"]
    labels = list(ct.index)
    out["rev_margin"] = C.revenue_margin_chart(labels, ct["revenue"], ct["ebitda"], ct["ebitda_margin"], charts_dir / "rev_margin.png",
                                              profit_label="EBITDA", margin_label="EBITDA margin", actual_count=sum(lab.endswith("A") for lab in labels),
                                              size=(4.0, 2.25))
    rel = r.relative_performance()
    if rel is not None:
        out["rel_perf"] = C.relative_performance_chart(rel, charts_dir / "rel_perf.png", stock_name=r.cfg.display_short,
                                                       index_name=r.cfg.index_name or r.cfg.index)
    # market slide: config charts first, then peer benchmarking fillers
    titles: dict[str, str] = {}
    i = 0
    for chart in r.cfg.market.charts[:4]:
        key = f"market_{i}"
        out[key] = C.config_chart(chart, charts_dir / f"{key}.png", size=(4.0, 2.2))
        titles[key] = chart.title
        i += 1
    tbl = r.comps.table
    fillers = []
    if tbl.empty:
        h = r.hist
        hl = [f"{int(y)}A" for y in h.index]
        ctx_labels = list(ct.index)
        fillers = [
            ("Revenue growth by year", "co_growth", lambda p: C.grouped_bars(
                ctx_labels, [("Revenue growth", [v * 100 if np.isfinite(v) else np.nan for v in ct["growth"]])], p, value_format="%",
                size=(4.0, 2.2), legend=False, highlight={lab: (NAVY if lab.endswith("A") else LIGHT_BLUE) for lab in ctx_labels})),
            ("Margin development", "co_margins", lambda p: C.line_chart(
                ctx_labels, [("EBITDA margin", [v * 100 for v in ct["ebitda_margin"]]), ("EBIT margin", [v * 100 for v in ct["ebit_margin"]]),
                             ("FCF margin", [v * 100 for v in ct["ufcf_margin"]])], p, value_format="%", size=(4.0, 2.2), show_y=True)),
            (f"Unlevered FCF and NOPAT ({r.units_label})", "co_fcf", lambda p: C.revenue_margin_chart(
                ctx_labels, ct["ufcf"], ct["nopat"], ct["ufcf_margin"], p, profit_label="NOPAT", margin_label="FCF margin",
                actual_count=sum(lab.endswith("A") for lab in ctx_labels), size=(4.0, 2.2), bar_label="Unlevered FCF")),
            ("Returns on capital", "co_returns", lambda p: C.grouped_bars(
                hl, [("ROIC", [v * 100 if np.isfinite(v) else np.nan for v in h["roic"]]), ("ROE", [v * 100 if np.isfinite(v) else np.nan for v in h["roe"]])],
                p, value_format="%", size=(4.0, 2.2))),
        ]
    else:
        own_g = float(r.hist["growth"].iloc[-1])
        own_m = float(r.hist["ebitda_margin"].iloc[-1])
        fillers = [
            ("Revenue growth vs peers (last fiscal year)", "peer_growth", lambda p: C.horizontal_bars(
                [r.cfg.display_short, *tbl["name"]], [own_g * 100, *[(v * 100 if v is not None and np.isfinite(v) else np.nan) for v in tbl["revenue_growth"]]],
                p, value_format="%", highlight_name=r.cfg.display_short, size=(4.0, 2.2))),
            ("EBITDA margin vs peers (lease-adjusted)", "peer_margin", lambda p: C.horizontal_bars(
                [r.cfg.display_short, *tbl["name"]], [own_m * 100, *[(v * 100 if v is not None and np.isfinite(v) else np.nan) for v in tbl["ebitda_margin"]]],
                p, value_format="%", highlight_name=r.cfg.display_short, size=(4.0, 2.2))),
            ("EV/EBITDA vs peers", "peer_ev_ebitda", lambda p: C.horizontal_bars(
                [r.cfg.display_short, *tbl["name"]], [r.comps.company.get("ev_ebitda"), *tbl["ev_ebitda"]], p, value_format="x",
                highlight_name=r.cfg.display_short, size=(4.0, 2.2))),
            ("P/E vs peers (trailing)", "peer_pe", lambda p: C.horizontal_bars(
                [r.cfg.display_short, *tbl["name"]], [r.comps.company.get("pe"), *tbl["pe"]], p, value_format="x",
                highlight_name=r.cfg.display_short, size=(4.0, 2.2))),
        ]
    for title, _key, fn in fillers:
        if i >= 4:
            break
        k = f"market_{i}"
        out[k] = fn(charts_dir / f"{k}.png")
        titles[k] = title
        i += 1
    out["_market_titles"] = titles  # type: ignore[assignment]
    # multiples panels
    stats = r.comps.stats
    for key in ("ev_ebitda", "ev_ebit", "pe"):
        if stats.empty:
            break
        vals = {"Peer avg": stats.loc["All|mean", key] if "All|mean" in stats.index else None,
                "Peer median": stats.loc["All|median", key] if "All|median" in stats.index else None,
                r.cfg.display_short: r.comps.company.get(key)}
        out[f"mult_{key}"] = C.multiples_panel(f"{MULTIPLE_LABELS[key]} (LTM)", vals, charts_dir / f"mult_{key}.png", highlight=r.cfg.display_short)
    # valuation slide
    out["football"] = C.football_field_chart(r.football, r.price, r.recommendation.target_price, charts_dir / "football.png", currency=r.price_currency)
    out["bridge_dcf"] = C.value_bars_chart([("Current", r.price), ("DCF base", r.dcf.value_per_share)], charts_dir / "bridge_dcf.png",
                                           size=(2.0, 2.3), colors=[GREY, NAVY])
    if len(r.scenarios) >= 2:
        sc_items = [("Current", r.price)] + [(s.name, s.value_per_share) for s in r.scenarios]
        sc_colors = [GREY] + [LIGHT_BLUE if s.name.lower() == "bear" else (GREEN if s.name.lower() == "bull" else NAVY) for s in r.scenarios]
        out["scenarios"] = C.value_bars_chart(sc_items, charts_dir / "scenarios.png", size=(2.0, 2.3), colors=sc_colors, decimals=0,
                                              title="DCF scenarios")
        out["scenarios_big"] = C.value_bars_chart(sc_items + ([("Weighted", r.scenario_weighted_value)] if r.scenario_weighted_value else []),
                                                  charts_dir / "scenarios_big.png", size=(4.4, 3.0), colors=sc_colors + [DARK_GREY], decimals=2,
                                                  title=f"Value per share by scenario ({r.price_currency})")
    items = [("Current", r.price)]
    for key, lab in (("ev_ebitda", "EV/EBITDA"), ("pe", "P/E")):
        iv = r.comps.implied.get(key)
        if iv is not None:
            items.append((lab, iv.per_share))
    if len(items) == 1:  # no peers: compare with the street's consensus target instead
        cons = (r.snapshot.price_targets or {}).get("mean")
        if cons:
            items.append(("Consensus TP", float(cons)))
    if len(items) > 1:
        out["bridge_mult"] = C.value_bars_chart(items, charts_dir / "bridge_mult.png", size=(2.0, 2.3), colors=[GREY, NAVY, LIGHT_BLUE])
    mh = getattr(r, "multiple_history", None)
    if mh is not None:
        for key, label in (("ev_ebitda", "EV/EBITDA"), ("pe", "P/E")):
            if key in mh.stats:
                out[f"hist_{key}"] = C.band_chart(mh.series[key], mh.stats[key], charts_dir / f"hist_{key}.png",
                                                  label=f"{r.cfg.display_short} trailing {label}")
    vd = [v for v in (getattr(r, "value_drivers", None) or []) if np.isfinite(v.value_down) and np.isfinite(v.value_up)]
    if len(vd) >= 2:
        out["tornado"] = C.tornado_chart(vd, vd[0].base, r.price, charts_dir / "tornado.png", currency=r.price_currency, size=(5.0, 4.4))
    reg = getattr(r.comps, "regression", None)
    if reg is not None and "growth_reg" in reg.coefficients and reg.actual_target:
        pts = [{"name": p["name"], "x": p["growth_reg"] * 100, "y": p["actual"]} for p in reg.fitted_peers if "growth_reg" in p]
        gt = reg.target_inputs.get("growth_reg")
        if pts and gt is not None:
            # the line is the fit at the target's own margin, so the chart is a plane cut through the target
            offset = reg.coefficients.get("ebitda_margin", 0.0) * reg.target_inputs.get("ebitda_margin", 0.0)
            out["mult_regression"] = C.scatter_regression_chart(
                pts, {"name": r.cfg.display_short, "x": gt * 100, "y": reg.actual_target}, charts_dir / "mult_regression.png",
                line=(reg.intercept + offset, reg.coefficients["growth_reg"] / 100), fitted_target=reg.fitted_target,
                x_label="Expected revenue growth, consensus FY1", line_label=f"Fit at {r.cfg.display_short}'s margin (R² {reg.r2:.2f})", size=(5.1, 2.75))
    return out


# =============================================================================
# entry point
# =============================================================================
def build_deck(result, out_path: str | Path, charts_dir: str | Path | None = None) -> Path:
    out_path = Path(out_path)
    charts_dir = Path(charts_dir) if charts_dir else out_path.parent / "charts"
    template = result.cfg.brand.template
    if template and Path(template).exists():
        prs = Presentation(template)
        _clear_template_slides(prs)
        template_used = True
    else:
        prs = Presentation()
        prs.slide_width = Inches(SW)
        prs.slide_height = Inches(SH)
        template_used = False
    ctx = DeckContext(result=result, template_used=template_used)
    ctx.charts = render_deck_charts(result, charts_dir)
    _cover(prs, ctx)
    _team(prs, ctx)
    _company(prs, ctx)
    _market(prs, ctx)
    _financials(prs, ctx)
    _valuation(prs, ctx)
    if result.cfg.appendix:
        _appendix_dcf(prs, ctx)
        _appendix_scenarios(prs, ctx)
        _appendix_peers(prs, ctx)
        _appendix_history(prs, ctx)
        _appendix_consensus(prs, ctx)
        _appendix_multiples(prs, ctx)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
