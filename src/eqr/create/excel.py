"""Excel model with live formulas (openpyxl).

Conventions (sell-side standard): blue = hard-coded input, black = formula,
green = link to another sheet. Every assumption sits in its own labelled cell on
the Inputs sheet and every forecast / valuation cell is a formula that
references it, so the analyst can flex the case directly in Excel.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .style import LIGHT_GREY, NAVY

BLUE = "0000FF"
BLACK = "000000"
GREEN = "008000"
FONT = "Arial"
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = "0.0%"
PCT2 = "0.00%"
MULT = '0.0"x"'
THIN = Side(style="thin", color="BFBFBF")
NAVY_FILL = PatternFill("solid", fgColor=NAVY)
GREY_FILL = PatternFill("solid", fgColor=LIGHT_GREY)
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")


class Sheet:
    """Small helper around an openpyxl worksheet."""

    def __init__(self, ws):
        self.ws = ws

    def cell(self, ref: str, value=None, *, fmt: str | None = None, color: str = BLACK, bold: bool = False, italic: bool = False,
             fill: PatternFill | None = None, align: str | None = None, size: int = 9, border: bool = False):
        c = self.ws[ref]
        c.value = value
        c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if align:
            c.alignment = Alignment(horizontal=align, vertical="center")
        if border:
            c.border = Border(bottom=THIN)
        return c

    def input(self, ref: str, value, fmt: str | None = None, **kw):
        return self.cell(ref, value, fmt=fmt, color=BLUE, **kw)

    def formula(self, ref: str, value: str, fmt: str | None = None, **kw):
        return self.cell(ref, value, fmt=fmt, color=BLACK, **kw)

    def link(self, ref: str, value: str, fmt: str | None = None, **kw):
        return self.cell(ref, value, fmt=fmt, color=GREEN, **kw)

    def header(self, ref: str, value, align: str = "center"):
        return self.cell(ref, value, bold=True, color="FFFFFF", fill=NAVY_FILL, align=align)

    def label(self, ref: str, value, bold: bool = False, italic: bool = False, indent: int = 0):
        c = self.cell(ref, value, bold=bold, italic=italic)
        if indent:
            c.alignment = Alignment(indent=indent)
        return c

    def widths(self, mapping: dict[str, float]):
        for col, w in mapping.items():
            self.ws.column_dimensions[col].width = w

    def title(self, text: str, sub: str = ""):
        self.cell("A1", text, bold=True, size=14, color=NAVY)
        if sub:
            self.cell("A2", sub, italic=True, size=9, color="595959")


def col(i: int) -> str:
    return get_column_letter(i)


def build_excel(result, out_path: str | Path) -> Path:
    r = result
    cfg = r.cfg
    out_path = Path(out_path)
    wb = Workbook()
    ws_cover = Sheet(wb.active)
    ws_cover.ws.title = "Cover"
    ws_in = Sheet(wb.create_sheet("Inputs"))
    ws_h = Sheet(wb.create_sheet("Historicals"))
    ws_f = Sheet(wb.create_sheet("Forecast"))
    ws_w = Sheet(wb.create_sheet("WACC"))
    ws_d = Sheet(wb.create_sheet("DCF"))
    ws_s = Sheet(wb.create_sheet("Scenarios"))
    ws_p = Sheet(wb.create_sheet("Peers"))
    ws_ff = Sheet(wb.create_sheet("Football field"))
    ws_c = Sheet(wb.create_sheet("Checks"))

    hist = r.hist
    fc = r.forecast[r.forecast.index != "TV"]
    years_h = [int(y) for y in hist.index]
    years_f = [int(y) for y in fc.index]
    n_f = len(years_f)
    units = r.units_label

    # ===================================================================== Inputs
    ws_in.title(f"{r.name} – model inputs", f"All figures in {units} unless stated. Blue = input, black = formula, green = link.")
    ws_in.widths({"A": 34, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14})
    ws_in.header("A4", "Market data", align="left")
    ws_in.header("B4", "Value")
    ws_in.header("C4", "Source / note", align="left")
    rows = [
        (f"Share price ({r.price_currency})", r.price, NUM2, "Yahoo Finance close"),
        (f"FX: {r.currency} per {r.price_currency}", r.fx_reporting_per_listing, "0.0000",
         "1.0 when the company reports in its listing currency" if not r.dual_currency else "Spot rate; per-share values are shown in the listing currency"),
        ("Shares outstanding (m)", r.shares_real, NUM1, "Yahoo Finance"),
        ("Net debt", r.net_debt, NUM, "Latest balance sheet" + (" (ex leases)" if cfg.assumptions.lease_treatment == "operating" else " (incl. leases)")),
        ("Minority interest", r.minorities, NUM, "Latest balance sheet"),
        ("Tax rate", r.drivers.tax_rate, PCT, r.drivers.tax_source or "Config"),
        ("Terminal growth", r.drivers.terminal_growth, PCT2, "Config"),
        ("Risk-free rate", r.wacc.risk_free, PCT2, "Config"),
        ("Equity risk premium", r.wacc.equity_risk_premium, PCT2, "Config"),
        ("Beta (adjusted)", r.wacc.beta_used, NUM2, r.wacc.beta_source),
        ("Size premium", r.wacc.size_premium, PCT2, r.wacc.size_premium_source or "Config"),
        ("Cost of debt (pre-tax)", r.wacc.cost_of_debt_pretax, PCT2, r.wacc.cost_of_debt_source),
        ("Gross debt for WACC weights", r.wacc.debt_value, NUM, "Latest balance sheet"),
        ("Mid-year convention (1 = yes)", 1 if cfg.assumptions.mid_year_convention else 0, "0", "Config"),
        ("Terminal method (1 = value driver)", 1 if r.dcf.terminal_method == "value_driver" else 0, "0",
         "0 = Gordon growth on FCF; 1 = NOPAT x (1 - g/RONIC) / (WACC - g)"),
        ("RONIC (value driver only)", r.dcf.ronic if r.dcf.ronic else r.wacc.wacc + 0.02, PCT2, "Return on new invested capital"),
        (f"Expected dividend per share ({r.price_currency})", r.recommendation.dps, NUM2, "Last year's cash dividend per share"),
        ("Roll-forward months", r.recommendation.horizon_months, "0", "12 = twelve-month target price; 0 = fair value today"),
        ("Target price rounding", cfg.recommendation.tp_rounding, NUM2, "Config"),
        ("Weight on DCF in fair value", cfg.recommendation.blend_dcf_weight if (cfg.recommendation.tp_method == "blend" and r.recommendation.multiples_value) else 1.0,
         PCT, "1 = pure DCF (tp_method: dcf); below 1 blends in the peer-multiple value"),
        (f"Peer-multiple value per share ({r.price_currency})", r.recommendation.multiples_value or 0.0, NUM2, "Median of the multiple-implied values (engine)"),
        ("BUY threshold (total return)", cfg.recommendation.buy_threshold, PCT, "Config"),
        ("SELL threshold (total return)", cfg.recommendation.sell_threshold, PCT, "Config"),
    ]
    IN: dict[str, str] = {}
    for i, (label, value, fmt, note) in enumerate(rows, start=5):
        ws_in.label(f"A{i}", label)
        ws_in.input(f"B{i}", value, fmt, fill=YELLOW_FILL if (label.startswith("Share price") or label in {"Terminal growth", "Beta (adjusted)"}) else None)
        ws_in.cell(f"C{i}", note, italic=True, color="595959")
        IN[label] = f"Inputs!$B${i}"
    IN["Share price"] = IN[f"Share price ({r.price_currency})"]
    IN["FX"] = IN[f"FX: {r.currency} per {r.price_currency}"]
    PX = f"({IN['Share price']}*{IN['FX']})"  # share price in the reporting currency
    SH = IN["Shares outstanding (m)"]
    FX = IN["FX"]
    # forecast drivers block
    r0 = 5 + len(rows) + 2
    ws_in.header(f"A{r0}", "Forecast drivers", align="left")
    for j, y in enumerate(years_f):
        ws_in.header(f"{col(2 + j)}{r0}", f"{y}E")
    driver_rows = [
        ("Revenue growth", r.drivers.revenue_growth, PCT),
        ("EBITDA margin", r.drivers.ebitda_margin, PCT),
        ("D&A % of revenue", [r.drivers.da_pct] * n_f, PCT),
        ("Capex % of revenue", [r.drivers.capex_at(j) for j in range(n_f)], PCT),
        ("Net working capital % of revenue", [r.drivers.nwc_pct] * n_f, PCT),
    ]
    DRV: dict[str, int] = {}
    for i, (label, values, fmt) in enumerate(driver_rows, start=r0 + 1):
        ws_in.label(f"A{i}", label)
        for j, v in enumerate(values):
            ws_in.input(f"{col(2 + j)}{i}", float(v), fmt, fill=YELLOW_FILL)
        DRV[label] = i
    note_row = r0 + len(driver_rows) + 2
    ws_in.cell(f"A{note_row}", "Notes: " + " ".join(r.drivers.notes), italic=True, color="595959")
    ws_in.cell(f"A{note_row + 1}", f"Lease treatment: {cfg.assumptions.lease_treatment} (lease rate {cfg.assumptions.lease_rate:.2%}). "
               f"Growth anchor: {r.drivers.growth_anchor}.", italic=True, color="595959")
    ws_in.ws.freeze_panes = "B5"

    # ================================================================ Historicals
    ws_h.title(f"{r.name} – historical financials ({units})", "Source: company reports via Yahoo Finance; lease adjustments per methodology.")
    ws_h.widths({"A": 34})
    ws_h.header("A4", units, align="left")
    for j, y in enumerate(years_h):
        ws_h.header(f"{col(2 + j)}4", f"{y}A")
        ws_h.ws.column_dimensions[col(2 + j)].width = 12
    h_rows = [
        ("Revenue", "revenue", NUM, "input"), ("Growth", None, PCT, "growth"),
        ("EBITDA (reported)", "ebitda_reported", NUM, "input"), ("Lease cost (rent)", "lease_cost", NUM, "input"),
        ("EBITDA", None, NUM, "ebitda"), ("EBITDA margin", None, PCT, "margin:EBITDA"),
        ("D&A (reported)", "da_reported", NUM, "input"), ("Right-of-use depreciation", "lease_principal", NUM, "input"),
        ("D&A", None, NUM, "da"), ("EBIT", None, NUM, "ebit"), ("EBIT margin", None, PCT, "margin:EBIT"),
        ("Tax rate (effective, used)", "tax_rate_used", PCT, "input"), ("NOPAT", None, NUM, "nopat"),
        ("Capex", "capex", NUM, "input"), ("Change in NWC", "nwc_change", NUM, "input"), ("Unlevered FCF", None, NUM, "ufcf"),
        ("FCF margin", None, PCT, "margin:Unlevered FCF"),
        ("Net income", "net_income", NUM, "input"), ("Diluted EPS", "eps", NUM2, "input"),
        ("Cash", "cash", NUM, "input"), ("Total debt (incl. leases)", "total_debt", NUM, "input"),
        ("Lease liabilities", "leases", NUM, "input"), ("Net debt (model basis)", "net_debt", NUM, "input"),
        ("Net debt / EBITDA", None, MULT, "ndebitda"), ("Interest expense (ex leases)", "interest_ex_leases", NUM, "input"),
    ]
    HR: dict[str, int] = {}
    for i, (label, key, fmt, kind) in enumerate(h_rows, start=5):
        HR[label] = i
        bold = label in {"Revenue", "EBITDA", "EBIT", "NOPAT", "Unlevered FCF"}
        ws_h.label(f"A{i}", label, bold=bold, italic=fmt == PCT)
        for j, _y in enumerate(years_h):
            c = f"{col(2 + j)}{i}"
            prev = f"{col(1 + j)}{i}"
            if kind == "input":
                v = hist[key].iloc[j]
                ws_h.input(c, None if pd.isna(v) else float(v), fmt)
            elif kind == "growth":
                prev_rev, cur_rev = f"{col(1 + j)}{HR['Revenue']}", f"{col(2 + j)}{HR['Revenue']}"
                ws_h.formula(c, f"=IF({prev_rev}=0,\"\",{cur_rev}/{prev_rev}-1)" if j else None, fmt, italic=True)
            elif kind == "ebitda":
                ws_h.formula(c, f"={col(2 + j)}{HR['EBITDA (reported)']}-{col(2 + j)}{HR['Lease cost (rent)']}", fmt, bold=True)
            elif kind == "da":
                ws_h.formula(c, f"=MAX({col(2 + j)}{HR['D&A (reported)']}-{col(2 + j)}{HR['Right-of-use depreciation']},0)", fmt)
            elif kind == "ebit":
                ws_h.formula(c, f"={col(2 + j)}{HR['EBITDA']}-{col(2 + j)}{HR['D&A']}", fmt, bold=True)
            elif kind == "nopat":
                ws_h.formula(c, f"={col(2 + j)}{HR['EBIT']}*(1-{col(2 + j)}{HR['Tax rate (effective, used)']})", fmt, bold=True)
            elif kind == "ufcf":
                ws_h.formula(c, f"={col(2 + j)}{HR['NOPAT']}+{col(2 + j)}{HR['D&A']}+{col(2 + j)}{HR['Capex']}+{col(2 + j)}{HR['Change in NWC']}", fmt, bold=True)
            elif kind.startswith("margin:"):
                base = kind.split(":", 1)[1]
                ws_h.formula(c, f"=IF({col(2 + j)}{HR['Revenue']}=0,\"\",{col(2 + j)}{HR[base]}/{col(2 + j)}{HR['Revenue']})", fmt, italic=True)
            elif kind == "ndebitda":
                ws_h.formula(c, f"=IF({col(2 + j)}{HR['EBITDA']}<=0,\"n.m.\",{col(2 + j)}{HR['Net debt (model basis)']}/{col(2 + j)}{HR['EBITDA']})", fmt)
    ws_h.ws.freeze_panes = "B5"

    # ==================================================================== Forecast
    ws_f.title(f"{r.name} – forecast ({units})", "Formulas reference the Inputs sheet; last actual year is linked from Historicals.")
    ws_f.widths({"A": 34, "B": 12})
    last_col_h = col(1 + len(years_h))
    ws_f.header("A4", units, align="left")
    ws_f.header("B4", f"{years_h[-1]}A")
    for j, y in enumerate(years_f):
        ws_f.header(f"{col(3 + j)}4", f"{y}E")
        ws_f.ws.column_dimensions[col(3 + j)].width = 12
    tv_col = col(3 + n_f)
    ws_f.header(f"{tv_col}4", "Terminal")
    ws_f.ws.column_dimensions[tv_col].width = 12
    FR = {"Revenue growth": 5, "Revenue": 6, "EBITDA margin": 7, "EBITDA": 8, "D&A % of revenue": 9, "D&A": 10, "EBIT": 11, "EBIT margin": 12,
          "Tax": 13, "NOPAT": 14, "Capex % of revenue": 15, "Capex": 16, "Net working capital % of revenue": 17, "Change in NWC": 18,
          "Unlevered FCF": 19, "FCF margin": 20}
    for label, i in FR.items():
        ws_f.label(f"A{i}", label, bold=label in {"Revenue", "EBITDA", "EBIT", "NOPAT", "Unlevered FCF"}, italic=("%" in label or "margin" in label))
    # last actual column (links)
    ws_f.link(f"B{FR['Revenue']}", f"=Historicals!{last_col_h}{HR['Revenue']}", NUM)
    ws_f.link(f"B{FR['EBITDA']}", f"=Historicals!{last_col_h}{HR['EBITDA']}", NUM)
    ws_f.formula(f"B{FR['EBITDA margin']}", f"=B{FR['EBITDA']}/B{FR['Revenue']}", PCT, italic=True)
    ws_f.link(f"B{FR['D&A']}", f"=Historicals!{last_col_h}{HR['D&A']}", NUM)
    ws_f.link(f"B{FR['EBIT']}", f"=Historicals!{last_col_h}{HR['EBIT']}", NUM)
    ws_f.link(f"B{FR['NOPAT']}", f"=Historicals!{last_col_h}{HR['NOPAT']}", NUM)
    ws_f.link(f"B{FR['Capex']}", f"=Historicals!{last_col_h}{HR['Capex']}", NUM)
    ws_f.link(f"B{FR['Change in NWC']}", f"=Historicals!{last_col_h}{HR['Change in NWC']}", NUM)
    ws_f.link(f"B{FR['Unlevered FCF']}", f"=Historicals!{last_col_h}{HR['Unlevered FCF']}", NUM)
    for j in range(n_f):
        c = col(3 + j)
        prev = col(2 + j)
        ic = col(2 + j)  # inputs sheet column for driver year j
        ws_f.link(f"{c}{FR['Revenue growth']}", f"=Inputs!{ic}{DRV['Revenue growth']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['Revenue']}", f"={prev}{FR['Revenue']}*(1+{c}{FR['Revenue growth']})", NUM, bold=True)
        ws_f.link(f"{c}{FR['EBITDA margin']}", f"=Inputs!{ic}{DRV['EBITDA margin']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['EBITDA']}", f"={c}{FR['Revenue']}*{c}{FR['EBITDA margin']}", NUM, bold=True)
        ws_f.link(f"{c}{FR['D&A % of revenue']}", f"=Inputs!{ic}{DRV['D&A % of revenue']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['D&A']}", f"={c}{FR['Revenue']}*{c}{FR['D&A % of revenue']}", NUM)
        ws_f.formula(f"{c}{FR['EBIT']}", f"={c}{FR['EBITDA']}-{c}{FR['D&A']}", NUM, bold=True)
        ws_f.formula(f"{c}{FR['EBIT margin']}", f"={c}{FR['EBIT']}/{c}{FR['Revenue']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['Tax']}", f"=-MAX({c}{FR['EBIT']},0)*{IN['Tax rate']}", NUM)
        ws_f.formula(f"{c}{FR['NOPAT']}", f"={c}{FR['EBIT']}+{c}{FR['Tax']}", NUM, bold=True)
        ws_f.link(f"{c}{FR['Capex % of revenue']}", f"=Inputs!{ic}{DRV['Capex % of revenue']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['Capex']}", f"=-{c}{FR['Revenue']}*{c}{FR['Capex % of revenue']}", NUM)
        ws_f.link(f"{c}{FR['Net working capital % of revenue']}", f"=Inputs!{ic}{DRV['Net working capital % of revenue']}", PCT, italic=True)
        ws_f.formula(f"{c}{FR['Change in NWC']}", f"=-{c}{FR['Net working capital % of revenue']}*({c}{FR['Revenue']}-{prev}{FR['Revenue']})", NUM)
        ws_f.formula(f"{c}{FR['Unlevered FCF']}", f"={c}{FR['NOPAT']}+{c}{FR['D&A']}+{c}{FR['Capex']}+{c}{FR['Change in NWC']}", NUM, bold=True)
        ws_f.formula(f"{c}{FR['FCF margin']}", f"={c}{FR['Unlevered FCF']}/{c}{FR['Revenue']}", PCT, italic=True)
    last_f = col(2 + n_f)
    t = tv_col
    ws_f.link(f"{t}{FR['Revenue growth']}", f"={IN['Terminal growth']}", PCT, italic=True)
    ws_f.formula(f"{t}{FR['Revenue']}", f"={last_f}{FR['Revenue']}*(1+{t}{FR['Revenue growth']})", NUM, bold=True)
    for label in ("EBITDA margin", "D&A % of revenue", "Capex % of revenue", "Net working capital % of revenue"):
        ws_f.formula(f"{t}{FR[label]}", f"={last_f}{FR[label]}", PCT, italic=True)
    ws_f.formula(f"{t}{FR['EBITDA']}", f"={t}{FR['Revenue']}*{t}{FR['EBITDA margin']}", NUM, bold=True)
    ws_f.formula(f"{t}{FR['D&A']}", f"={t}{FR['Revenue']}*{t}{FR['D&A % of revenue']}", NUM)
    ws_f.formula(f"{t}{FR['EBIT']}", f"={t}{FR['EBITDA']}-{t}{FR['D&A']}", NUM, bold=True)
    ws_f.formula(f"{t}{FR['EBIT margin']}", f"={t}{FR['EBIT']}/{t}{FR['Revenue']}", PCT, italic=True)
    ws_f.formula(f"{t}{FR['Tax']}", f"=-MAX({t}{FR['EBIT']},0)*{IN['Tax rate']}", NUM)
    ws_f.formula(f"{t}{FR['NOPAT']}", f"={t}{FR['EBIT']}+{t}{FR['Tax']}", NUM, bold=True)
    ws_f.formula(f"{t}{FR['Capex']}", f"=-{t}{FR['Revenue']}*{t}{FR['Capex % of revenue']}", NUM)
    ws_f.formula(f"{t}{FR['Change in NWC']}", f"=-{t}{FR['Net working capital % of revenue']}*({t}{FR['Revenue']}-{last_f}{FR['Revenue']})", NUM)
    ws_f.formula(f"{t}{FR['Unlevered FCF']}", f"={t}{FR['NOPAT']}+{t}{FR['D&A']}+{t}{FR['Capex']}+{t}{FR['Change in NWC']}", NUM, bold=True)
    ws_f.formula(f"{t}{FR['FCF margin']}", f"={t}{FR['Unlevered FCF']}/{t}{FR['Revenue']}", PCT, italic=True)
    ws_f.ws.freeze_panes = "C5"

    # ======================================================================== WACC
    ws_w.title("Cost of capital", "CAPM cost of equity, after-tax cost of debt, market-value weights.")
    ws_w.widths({"A": 34, "B": 14, "C": 40})
    w_rows = [
        ("Risk-free rate", f"={IN['Risk-free rate']}", PCT2, "link"),
        ("Beta (adjusted)", f"={IN['Beta (adjusted)']}", NUM2, "link"),
        ("Equity risk premium", f"={IN['Equity risk premium']}", PCT2, "link"),
        ("Size premium", f"={IN['Size premium']}", PCT2, "link"),
        ("Cost of equity", "=B5+B6*B7+B8", PCT2, "formula"),
        ("Cost of debt (pre-tax)", f"={IN['Cost of debt (pre-tax)']}", PCT2, "link"),
        ("Tax rate", f"={IN['Tax rate']}", PCT, "link"),
        ("Cost of debt (after tax)", "=B10*(1-B11)", PCT2, "formula"),
        ("Equity value (market cap)", f"={PX}*{SH}", NUM, "formula"),
        ("Debt value", f"={IN['Gross debt for WACC weights']}", NUM, "link"),
        ("Weight of equity", "=B13/(B13+B14)", PCT, "formula"),
        ("Weight of debt", "=B14/(B13+B14)", PCT, "formula"),
        ("WACC", "=B15*B9+B16*B12", PCT2, "formula"),
    ]
    ws_w.header("A4", "Component", align="left")
    ws_w.header("B4", "Value")
    for i, (label, f, fmt, kind) in enumerate(w_rows, start=5):
        ws_w.label(f"A{i}", label, bold=label in {"Cost of equity", "WACC"})
        (ws_w.link if kind == "link" else ws_w.formula)(f"B{i}", f, fmt, bold=label in {"Cost of equity", "WACC"})
    ws_w.cell("C6", r.wacc.beta_source, italic=True, color="595959")
    WACC_REF = "WACC!$B$17"

    # ========================================================================= DCF
    ws_d.title(f"{r.name} – DCF valuation ({units})", "Unlevered FCF discounted at WACC; Gordon growth terminal value.")
    ws_d.widths({"A": 34})
    ws_d.header("A4", "Year", align="left")
    for j, y in enumerate(years_f):
        ws_d.header(f"{col(2 + j)}4", f"{y}E")
        ws_d.ws.column_dimensions[col(2 + j)].width = 12
    ws_d.label("A5", "Unlevered FCF", bold=True)
    ws_d.label("A6", "Discount period (years)")
    ws_d.label("A7", "Discount factor")
    ws_d.label("A8", "PV of FCF", bold=True)
    for j in range(n_f):
        c = col(2 + j)
        ws_d.link(f"{c}5", f"=Forecast!{col(3 + j)}{FR['Unlevered FCF']}", NUM, bold=True)
        ws_d.formula(f"{c}6", f"={j + 1}-0.5*{IN['Mid-year convention (1 = yes)']}", NUM1)
        ws_d.formula(f"{c}7", f"=1/(1+{WACC_REF})^{c}6", "0.000")
        ws_d.formula(f"{c}8", f"={c}5*{c}7", NUM, bold=True)
    METHOD, RONIC = IN["Terminal method (1 = value driver)"], IN["RONIC (value driver only)"]
    NOPAT_N = f"Forecast!${last_f}${FR['NOPAT']}"
    MONTHS, DPS = IN["Roll-forward months"], IN[f"Expected dividend per share ({r.price_currency})"]

    def tcf(g: str, fcf_n: str = f"${col(1 + n_f)}$5", nopat_n: str = NOPAT_N) -> str:
        """Excel expression for the terminal-year cash flow given a growth-rate reference."""
        return f"IF({METHOD}=1,{nopat_n}*(1+{g})*(1-{g}/{RONIC}),{fcf_n}*(1+{g}))"

    d_rows = [
        ("Sum of PV of FCF", f"=SUM(B8:{col(1 + n_f)}8)", NUM, True),
        ("Terminal growth", f"={IN['Terminal growth']}", PCT2, False),
        ("WACC", f"={WACC_REF}", PCT2, False),
        ("Terminal cash flow (year N+1)", f"={tcf('B11')}", NUM, False),
        ("Terminal value", "=B13/(B12-B11)", NUM, True),
        ("PV of terminal value", f"=B14/(1+B12)^{n_f}", NUM, True),
        ("Enterprise value", "=B10+B15", NUM, True),
        ("Less: net debt", f"=-{IN['Net debt']}", NUM, False),
        ("Less: minority interest", f"=-{IN['Minority interest']}", NUM, False),
        ("Equity value", "=B16+B17+B18", NUM, True),
        ("Shares outstanding (m)", f"={SH}", NUM1, False),
        (f"Fair value per share, DCF ({r.price_currency})", f"=B19/B20/{FX}", NUM2, True),
        (f"Current share price ({r.price_currency})", f"={IN['Share price']}", NUM2, False),
        ("Upside / (downside) to fair value", "=B21/B22-1", PCT, True),
        ("Terminal value as % of EV", "=B15/B16", PCT, False),
        ("Implied exit EV/EBITDA", f"=B14/Forecast!{last_f}{FR['EBITDA']}", MULT, False),
        ("Fair value used for the target price", f"={IN['Weight on DCF in fair value']}*B21+(1-{IN['Weight on DCF in fair value']})*"
                                                 f"{IN[f'Peer-multiple value per share ({r.price_currency})']}", NUM2, False),
        ("Cost of equity", "=WACC!$B$9", PCT2, False),
        (f"Expected dividend per share ({r.price_currency})", f"=IF({MONTHS}>0,{DPS},0)", NUM2, False),
        ("Target price before rounding", f"=IF(B26<=0,0,B26*(1+B27)^({MONTHS}/12)-B28)", NUM2, False),
        (f"Target price ({r.price_currency})", f"=IF({IN['Target price rounding']}>0,MROUND(B29,{IN['Target price rounding']}),B29)", NUM2, True),
        ("Upside to target price", "=B30/B22-1", PCT, True),
        ("Expected total return (incl. dividend)", "=(B30+B28)/B22-1", PCT, True),
        ("Rating", f"=IF(B26<=0,\"NOT RATED\",IF(B32>={IN['BUY threshold (total return)']},\"BUY\",IF(B32<={IN['SELL threshold (total return)']},\"SELL\",\"HOLD\")))", None, True),
    ]
    for i, (label, f, fmt, bold) in enumerate(d_rows, start=10):
        ws_d.label(f"A{i}", label, bold=bold)
        (ws_d.link if f.startswith("=Inputs") or f.startswith("=WACC") or f.startswith("=-Inputs") else ws_d.formula)(f"B{i}", f, fmt, bold=bold)
    ws_d.ws["B21"].fill = YELLOW_FILL
    ws_d.ws["B30"].fill = YELLOW_FILL
    # ---- sensitivity (live formulas)
    s0 = 37
    ws_d.cell(f"A{s0 - 1}", "Sensitivity: value per share (rows = terminal growth, columns = WACC)", bold=True, color=NAVY)
    waccs, growths = r.dcf.sensitivity_waccs, r.dcf.sensitivity_growths
    k = len(waccs) // 2
    ws_d.header(f"A{s0}", "g \\ WACC")
    for j, _w in enumerate(waccs):
        c = col(2 + j)
        ws_d.formula(f"{c}{s0}", f"={WACC_REF}+({j - k})*{cfg.recommendation.sensitivity_wacc_step}", PCT2, bold=True, fill=GREY_FILL, align="center")
    fcf_range = f"$B$5:${col(1 + n_f)}$5"
    mid = IN["Mid-year convention (1 = yes)"]
    for i, _g in enumerate(growths):
        rr = s0 + 1 + i
        ws_d.formula(f"A{rr}", f"={IN['Terminal growth']}+({i - k})*{cfg.recommendation.sensitivity_growth_step}", PCT2, bold=True, fill=GREY_FILL, align="center")
        for j in range(len(waccs)):
            c = col(2 + j)
            w_ref, g_ref = f"{c}${s0}", f"$A{rr}"
            formula = (f"=IF({w_ref}<={g_ref},\"n.m.\",(NPV({w_ref},{fcf_range})*(1+{w_ref})^(0.5*{mid})"
                       f"+{tcf(g_ref)}/({w_ref}-{g_ref})/(1+{w_ref})^{n_f}-{IN['Net debt']}-{IN['Minority interest']})/{SH}/{FX})")
            ws_d.formula(f"{c}{rr}", formula, NUM1, bold=(i == k and j == k), fill=YELLOW_FILL if (i == k and j == k) else None, align="center")
    ws_d.ws.freeze_panes = "B5"

    # ======================================================================= Peers
    ws_p.title("Peer multiples", "Multiples computed from last reported fiscal year; EV on the same lease basis as the target company.")
    ws_p.widths({"A": 24, "B": 12, "C": 18, "D": 14, "E": 14, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12, "L": 12})
    headers = ["Company", "Ticker", "Group", f"Market cap ({units})", f"EV ({units})", "Rev. growth", "EBITDA margin", "EV/Sales", "EV/EBITDA", "EV/EBIT", "P/E", "P/E (fwd)"]
    for j, hname in enumerate(headers):
        ws_p.header(f"{col(1 + j)}4", hname, align="left" if j < 3 else "center")
    tbl = r.comps.table
    keys = ["name", "ticker", "group", "market_cap", "ev", "revenue_growth", "ebitda_margin", "ev_sales", "ev_ebitda", "ev_ebit", "pe", "fwd_pe"]
    fmts = [None, None, None, NUM, NUM, PCT, PCT, MULT, MULT, MULT, MULT, MULT]
    first_peer, last_peer = 5, 4 + len(tbl)
    for i, (_, row) in enumerate(tbl.iterrows(), start=5):
        for j, (key, fmt) in enumerate(zip(keys, fmts)):
            v = row.get(key)
            if isinstance(v, float) and np.isnan(v):
                v = None
            ws_p.input(f"{col(1 + j)}{i}", v, fmt) if j >= 3 else ws_p.label(f"{col(1 + j)}{i}", v)
    stat_start = last_peer + 2
    for s_i, (label, fn) in enumerate([("Average", "AVERAGE"), ("Median", "MEDIAN"), ("25th percentile", "QUARTILE"), ("75th percentile", "QUARTILE")]):
        rr = stat_start + s_i
        ws_p.label(f"A{rr}", label, bold=True)
        for j in range(5, 12):
            c = col(1 + j)
            rng = f"{c}{first_peer}:{c}{last_peer}"
            if fn == "QUARTILE":
                q = 1 if "25" in label else 3
                f = f"=IFERROR(QUARTILE({rng},{q}),\"\")"
            else:
                f = f"=IFERROR({fn}({rng}),\"\")"
            ws_p.formula(f"{c}{rr}", f, fmts[j], bold=True, border=(s_i == 0))
    med_row = stat_start + 1
    # company block
    cb = stat_start + 6
    ws_p.cell(f"A{cb}", f"{cfg.display_short} – current multiples and implied value", bold=True, color=NAVY)
    last_h = last_col_h
    ws_p.label(f"A{cb + 1}", "Market cap")
    ws_p.formula(f"B{cb + 1}", f"={PX}*{SH}", NUM)
    ws_p.label(f"A{cb + 2}", "Enterprise value")
    ws_p.formula(f"B{cb + 2}", f"=B{cb + 1}+{IN['Net debt']}+{IN['Minority interest']}", NUM)
    ws_p.label(f"A{cb + 3}", "LTM revenue")
    ws_p.link(f"B{cb + 3}", f"=Historicals!{last_h}{HR['Revenue']}", NUM)
    ws_p.label(f"A{cb + 4}", "LTM EBITDA")
    ws_p.link(f"B{cb + 4}", f"=Historicals!{last_h}{HR['EBITDA']}", NUM)
    ws_p.label(f"A{cb + 5}", "LTM EBIT")
    ws_p.link(f"B{cb + 5}", f"=Historicals!{last_h}{HR['EBIT']}", NUM)
    ws_p.label(f"A{cb + 6}", "LTM EPS")
    ws_p.link(f"B{cb + 6}", f"=Historicals!{last_h}{HR['Diluted EPS']}", NUM2)
    ws_p.header(f"D{cb}", "Multiple", align="left")
    ws_p.header(f"E{cb}", cfg.display_short)
    ws_p.header(f"F{cb}", "Peer median")
    ws_p.header(f"G{cb}", f"Implied value/share ({r.price_currency})")
    ws_p.header(f"H{cb}", "vs price")
    implied_rows = [("EV/Sales", "H", f"B{cb + 3}", True), ("EV/EBITDA", "I", f"B{cb + 4}", True), ("EV/EBIT", "J", f"B{cb + 5}", True), ("P/E", "K", f"B{cb + 6}", False)]
    IMPL: dict[str, int] = {}
    for i, (label, mcol, metric, ev_based) in enumerate(implied_rows, start=cb + 1):
        IMPL[label] = i
        ws_p.label(f"D{i}", label)
        if ev_based:
            ws_p.formula(f"E{i}", f"=IF({metric}<=0,\"n.m.\",B{cb + 2}/{metric})", MULT)
            ws_p.formula(f"G{i}", f"=IF({metric}<=0,\"\",({mcol}{med_row}*{metric}-{IN['Net debt']}-{IN['Minority interest']})/{SH}/{FX})", NUM2)
        else:
            ws_p.formula(f"E{i}", f"=IF({metric}<=0,\"n.m.\",{PX}/{metric})", MULT)
            ws_p.formula(f"G{i}", f"=IF({metric}<=0,\"\",{mcol}{med_row}*{metric}/{FX})", NUM2)
        ws_p.formula(f"F{i}", f"={mcol}{med_row}", MULT)
        ws_p.formula(f"H{i}", f"=IF(G{i}=\"\",\"\",G{i}/{IN['Share price']}-1)", PCT)
    ws_p.ws.freeze_panes = "B5"

    # ============================================================== Football field
    ws_ff.title("Valuation summary", "Ranges feed the football-field chart; DCF range = inner 3x3 of the sensitivity grid.")
    ws_ff.widths({"A": 34, "B": 12, "C": 12, "D": 12})
    for j, hname in enumerate(["Method", "Low", "High", "Range"]):
        ws_ff.header(f"{col(1 + j)}4", hname, align="left" if j == 0 else "center")
    sens_inner = f"DCF!{col(2 + k - 1)}{s0 + k}:{col(2 + k + 1)}{s0 + k + 2}"
    ff_rows = [
        ("52-week range", r.snapshot.info.get("fiftyTwoWeekLow"), r.snapshot.info.get("fiftyTwoWeekHigh"), True),
        ("DCF (WACC ±0.5%, g ±0.25%)", f"=MIN({sens_inner})", f"=MAX({sens_inner})", False),
    ]
    p25_row, p75_row = stat_start + 2, stat_start + 3
    for label, mcol, metric_ref, ev_based in implied_rows[1:]:
        if ev_based:
            lo = f"=(Peers!{mcol}{p25_row}*Peers!{metric_ref}-{IN['Net debt']}-{IN['Minority interest']})/{SH}/{FX}"
            hi = f"=(Peers!{mcol}{p75_row}*Peers!{metric_ref}-{IN['Net debt']}-{IN['Minority interest']})/{SH}/{FX}"
        else:
            lo = f"=Peers!{mcol}{p25_row}*Peers!{metric_ref}/{FX}"
            hi = f"=Peers!{mcol}{p75_row}*Peers!{metric_ref}/{FX}"
        ff_rows.append((f"{label} (peer 25th–75th pct)", lo, hi, False))
    for i, (label, lo, hi, is_input) in enumerate(ff_rows, start=5):
        ws_ff.label(f"A{i}", label)
        (ws_ff.input if is_input else ws_ff.formula)(f"B{i}", lo, NUM2)
        (ws_ff.input if is_input else ws_ff.formula)(f"C{i}", hi, NUM2)
        ws_ff.formula(f"D{i}", f"=IFERROR(C{i}-B{i},\"\")", NUM2)
    last_ff = 4 + len(ff_rows)
    ws_ff.label(f"A{last_ff + 2}", f"Current share price ({r.price_currency})")
    ws_ff.link(f"B{last_ff + 2}", f"={IN['Share price']}", NUM2)
    ws_ff.label(f"A{last_ff + 3}", f"Target price ({r.price_currency})")
    ws_ff.link(f"B{last_ff + 3}", "=DCF!$B$30", NUM2, fill=YELLOW_FILL)
    ws_ff.label(f"A{last_ff + 4}", "Rating")
    ws_ff.link(f"B{last_ff + 4}", "=DCF!$B$33", None, bold=True)
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Football field (value per share)"
    chart.style = 10
    data = Reference(ws_ff.ws, min_col=2, min_row=4, max_col=2, max_row=last_ff)
    rng = Reference(ws_ff.ws, min_col=4, min_row=4, max_col=4, max_row=last_ff)
    cats = Reference(ws_ff.ws, min_col=1, min_row=5, max_row=last_ff)
    chart.add_data(data, titles_from_data=True)
    chart.add_data(rng, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.noFill = True
    chart.series[0].graphicalProperties.line.noFill = True
    chart.series[1].graphicalProperties.solidFill = NAVY
    chart.legend = None
    chart.height, chart.width = 8, 18
    ws_ff.ws.add_chart(chart, "F4")

    # =================================================================== Scenarios
    ws_s.title("Scenario analysis", "Shifts are additive to the base case (Inputs sheet). Edit the blue cells; every scenario re-runs the full forecast.")
    ws_s.widths({"A": 30})
    for j in range(2, 3 + n_f + 2):
        ws_s.ws.column_dimensions[col(j)].width = 13
    sc_headers = ["Scenario", "Probability", "Growth shift", "Margin shift", "WACC shift", "Terminal g shift", "WACC", "Terminal growth",
                  f"Value per share ({r.price_currency})", "vs price"]
    for j, hname in enumerate(sc_headers):
        ws_s.header(f"{col(1 + j)}4", hname, align="left" if j == 0 else "center")
    sc_list = list(cfg.scenarios)
    n_sc = len(sc_list)
    block0 = 5 + n_sc + 4
    block_h = 15
    mid_ref = IN["Mid-year convention (1 = yes)"]
    for si, sc in enumerate(sc_list):
        row = 5 + si
        b = block0 + si * block_h  # first row of this scenario's forecast block
        ws_s.label(f"A{row}", sc.name, bold=True)
        ws_s.input(f"B{row}", float(sc.probability), PCT)
        ws_s.input(f"C{row}", float(sc.growth_shift), PCT)
        ws_s.input(f"D{row}", float(sc.margin_shift), PCT)
        ws_s.input(f"E{row}", float(sc.wacc_shift), PCT2)
        ws_s.input(f"F{row}", float(sc.terminal_growth_shift), PCT2)
        ws_s.formula(f"G{row}", f"={WACC_REF}+E{row}", PCT2)
        ws_s.formula(f"H{row}", f"={IN['Terminal growth']}+F{row}", PCT2)
        ws_s.formula(f"I{row}", f"=B{b + 13}", NUM2, bold=True)
        ws_s.formula(f"J{row}", f"=I{row}/{IN['Share price']}-1", PCT)
        # ---- forecast block
        ws_s.cell(f"A{b}", f"{sc.name} case ({units})", bold=True, color=NAVY)
        ws_s.link(f"B{b}", "=Forecast!B4", None, bold=True, align="center")
        for j, y in enumerate(years_f):
            ws_s.header(f"{col(3 + j)}{b}", f"{y}E")
        labels = ["Revenue growth", "Revenue", "EBITDA margin", "EBITDA", "D&A", "EBIT", "Tax", "NOPAT", "Capex", "Change in NWC", "Unlevered FCF"]
        R = {lab: b + 1 + i for i, lab in enumerate(labels)}
        for lab in labels:
            ws_s.label(f"A{R[lab]}", lab, bold=lab in {"Revenue", "EBITDA", "Unlevered FCF"}, italic=lab in {"Revenue growth", "EBITDA margin"})
        ws_s.link(f"B{R['Revenue']}", f"=Forecast!B{FR['Revenue']}", NUM)
        for j in range(n_f):
            c, prev, ic = col(3 + j), col(2 + j), col(2 + j)
            ws_s.formula(f"{c}{R['Revenue growth']}", f"=Inputs!{ic}{DRV['Revenue growth']}+$C${row}", PCT, italic=True)
            ws_s.formula(f"{c}{R['Revenue']}", f"={prev}{R['Revenue']}*(1+{c}{R['Revenue growth']})", NUM, bold=True)
            ws_s.formula(f"{c}{R['EBITDA margin']}", f"=Inputs!{ic}{DRV['EBITDA margin']}+$D${row}", PCT, italic=True)
            ws_s.formula(f"{c}{R['EBITDA']}", f"={c}{R['Revenue']}*{c}{R['EBITDA margin']}", NUM, bold=True)
            ws_s.formula(f"{c}{R['D&A']}", f"={c}{R['Revenue']}*Inputs!{ic}{DRV['D&A % of revenue']}", NUM)
            ws_s.formula(f"{c}{R['EBIT']}", f"={c}{R['EBITDA']}-{c}{R['D&A']}", NUM)
            ws_s.formula(f"{c}{R['Tax']}", f"=-MAX({c}{R['EBIT']},0)*{IN['Tax rate']}", NUM)
            ws_s.formula(f"{c}{R['NOPAT']}", f"={c}{R['EBIT']}+{c}{R['Tax']}", NUM)
            ws_s.formula(f"{c}{R['Capex']}", f"=-{c}{R['Revenue']}*Inputs!{ic}{DRV['Capex % of revenue']}", NUM)
            ws_s.formula(f"{c}{R['Change in NWC']}", f"=-Inputs!{ic}{DRV['Net working capital % of revenue']}*({c}{R['Revenue']}-{prev}{R['Revenue']})", NUM)
            ws_s.formula(f"{c}{R['Unlevered FCF']}", f"={c}{R['NOPAT']}+{c}{R['D&A']}+{c}{R['Capex']}+{c}{R['Change in NWC']}", NUM, bold=True)
        last_c = col(2 + n_f)
        r_fcf, r_nopat = R["Unlevered FCF"], R["NOPAT"]
        fcf_rng = f"$C${r_fcf}:${last_c}${r_fcf}"
        tcf_sc = tcf(f"$H${row}", fcf_n=f"${last_c}${r_fcf}", nopat_n=f"${last_c}${r_nopat}")
        ws_s.label(f"A{b + 12}", "Enterprise value")
        ws_s.formula(f"B{b + 12}", f"=NPV($G${row},{fcf_rng})*(1+$G${row})^(0.5*{mid_ref})+{tcf_sc}/($G${row}-$H${row})/(1+$G${row})^{n_f}", NUM)
        ws_s.label(f"A{b + 13}", f"Value per share ({r.price_currency})", bold=True)
        ws_s.formula(f"B{b + 13}", f"=(B{b + 12}-{IN['Net debt']}-{IN['Minority interest']})/{SH}/{FX}", NUM2, bold=True)
    wrow = 5 + n_sc
    ws_s.label(f"A{wrow}", "Probability-weighted", bold=True)
    ws_s.formula(f"B{wrow}", f"=SUM(B5:B{wrow - 1})", PCT, bold=True)
    ws_s.formula(f"I{wrow}", f"=IFERROR(SUMPRODUCT(B5:B{wrow - 1},I5:I{wrow - 1})/B{wrow},\"\")", NUM2, bold=True, fill=YELLOW_FILL)
    ws_s.formula(f"J{wrow}", f"=IFERROR(I{wrow}/{IN['Share price']}-1,\"\")", PCT, bold=True)
    SC_WEIGHTED = f"$I${wrow}"
    base_rows = [5 + i for i, s_ in enumerate(sc_list) if abs(s_.growth_shift) + abs(s_.margin_shift) + abs(s_.wacc_shift) + abs(s_.terminal_growth_shift) == 0]
    ws_s.ws.freeze_panes = "B5"

    # ====================================================================== Checks
    ws_c.title("Model integrity checks", "Every check must read TRUE. A FALSE means an input or formula has been broken.")
    ws_c.widths({"A": 62, "B": 12, "C": 60})
    ws_c.header("A4", "Check", align="left")
    ws_c.header("B4", "Result")
    ws_c.header("C4", "What it protects against", align="left")
    k_centre = len(r.dcf.sensitivity_waccs) // 2
    centre = f"DCF!{col(2 + k_centre)}{s0 + 1 + k_centre}"
    checks = [
        ("WACC exceeds terminal growth", f"={WACC_REF}>{IN['Terminal growth']}", "A perpetuity needs WACC > g"),
        ("Enterprise value = PV of FCF + PV of terminal value", "=ABS(DCF!B16-(DCF!B10+DCF!B15))<0.01", "Broken bridge formulas"),
        ("Equity value = EV - net debt - minorities", f"=ABS(DCF!B19-(DCF!B16-{IN['Net debt']}-{IN['Minority interest']}))<0.01", "Sign errors in the bridge"),
        ("Sensitivity centre equals the DCF fair value", f"=ABS({centre}-DCF!B21)<0.01", "Sensitivity grid out of sync with the model"),
        ("Forecast starts from the last reported revenue", f"=Forecast!B{FR['Revenue']}=Historicals!{last_col_h}{HR['Revenue']}", "Stale links"),
        ("Capital structure weights sum to 100%", "=ABS(WACC!B15+WACC!B16-1)<0.000001", "WACC weights"),
        ("Scenario probabilities sum to 100%", f"=ABS(Scenarios!B{wrow}-1)<0.000001", "Probability-weighted value"),
        ("Share price, shares and FX are positive", f"=AND({IN['Share price']}>0,{SH}>0,{FX}>0)", "Missing market data"),
        ("Terminal value below 85% of enterprise value", "=DCF!B24<0.85", "Valuation resting almost entirely on the perpetuity"),
        ("No error values in Forecast, DCF or Scenarios",
         f"=SUMPRODUCT(--ISERROR(Forecast!B5:{tv_col}20))+SUMPRODUCT(--ISERROR(DCF!B5:{col(1 + n_f)}33))+SUMPRODUCT(--ISERROR(Scenarios!B5:J{wrow}))=0",
         "#REF!, #DIV/0! and friends"),
    ]
    if base_rows:
        checks.insert(4, ("Base scenario reproduces the DCF fair value", f"=ABS(Scenarios!I{base_rows[0]}-DCF!B21)<0.01", "Scenario blocks drifting from the forecast"))
    for i, (label, f, why) in enumerate(checks, start=5):
        ws_c.label(f"A{i}", label)
        ws_c.formula(f"B{i}", f, None, bold=True, align="center")
        ws_c.cell(f"C{i}", why, italic=True, color="595959")
    status_row = 5 + len(checks) + 1
    ws_c.label(f"A{status_row}", "Overall status", bold=True)
    ws_c.formula(f"B{status_row}", f"=IF(AND(B5:B{status_row - 2}),\"ALL OK\",\"CHECK\")", None, bold=True, align="center", fill=YELLOW_FILL)
    CHECK_STATUS = f"$B${status_row}"

    # ======================================================================= Cover
    ws_cover.title(f"{r.name} ({cfg.ticker}) – valuation model", f"Generated by Equity Research Engine on {r.as_of.isoformat()}")
    ws_cover.widths({"A": 30, "B": 18, "C": 60})
    cover_rows = [
        (f"Share price ({r.price_currency})", f"={IN['Share price']}", NUM2), (f"DCF fair value per share ({r.price_currency})", "=DCF!B21", NUM2),
        (f"Probability-weighted scenario value ({r.price_currency})", f"=Scenarios!{SC_WEIGHTED}", NUM2),
        (f"Target price ({r.price_currency})", "=DCF!B30", NUM2), ("Upside to target price", "=DCF!B31", PCT),
        ("Expected total return", "=DCF!B32", PCT), ("Rating", "=DCF!B33", None), ("Model checks", f"=Checks!{CHECK_STATUS}", None),
        ("WACC", f"={WACC_REF}", PCT2), ("Terminal growth", f"={IN['Terminal growth']}", PCT2),
        ("EV/EBITDA (LTM)", f"=Peers!E{IMPL['EV/EBITDA']}", MULT), ("Peer median EV/EBITDA", f"=Peers!F{IMPL['EV/EBITDA']}", MULT),
    ]
    for i, (label, f, fmt) in enumerate(cover_rows, start=4):
        ws_cover.label(f"A{i}", label, bold=True)
        ws_cover.link(f"B{i}", f, fmt)
    legend_row = 4 + len(cover_rows) + 2
    ws_cover.cell(f"A{legend_row}", "Legend", bold=True, color=NAVY)
    ws_cover.cell(f"A{legend_row + 1}", "Blue text = hard-coded input (edit these)", color=BLUE)
    ws_cover.cell(f"A{legend_row + 2}", "Black text = formula", color=BLACK)
    ws_cover.cell(f"A{legend_row + 3}", "Green text = link to another sheet", color=GREEN)
    ws_cover.cell(f"A{legend_row + 4}", "Yellow fill = key assumption", fill=YELLOW_FILL)
    ws_cover.cell(f"A{legend_row + 6}", "Sheets: Inputs -> Historicals -> Forecast -> WACC -> DCF -> Peers -> Football field", italic=True, color="595959")
    ws_cover.cell(f"A{legend_row + 7}", f"Lease treatment: {cfg.assumptions.lease_treatment}. Sources: {cfg.sources_note}", italic=True, color="595959")
    if r.dual_currency:
        ws_cover.cell(f"A{legend_row + 8}", f"Accounts in {r.currency}; share price and per-share values in {r.price_currency} "
                      f"(FX {r.fx_reporting_per_listing:.4f} {r.currency} per {r.price_currency} on the Inputs sheet).", italic=True, color="595959")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    try:
        recalc_with_excel(out_path)
    except Exception:
        pass
    return out_path


def recalc_with_excel(path: str | Path) -> dict:
    """Recalculate with Excel (Windows COM) so cached values exist; report error cells."""
    import sys

    if not sys.platform.startswith("win"):
        return {"status": "skipped"}
    import pythoncom  # type: ignore
    import win32com.client  # type: ignore

    pythoncom.CoInitialize()
    app = win32com.client.Dispatch("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    errors: list[str] = []
    wb = app.Workbooks.Open(str(Path(path).resolve()))
    try:
        app.CalculateFull()
        for ws in wb.Worksheets:
            try:
                rng = ws.UsedRange.SpecialCells(-4123, 16)  # xlCellTypeFormulas, xlErrors
                for c in rng:
                    errors.append(f"{ws.Name}!{c.Address(False, False)}={c.Text}")
            except Exception:
                pass
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)
        app.Quit()
    return {"status": "errors_found" if errors else "success", "errors": errors}
